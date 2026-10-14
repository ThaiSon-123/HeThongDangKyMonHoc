"""Pha 9: báo cáo đào tạo (PDT-14), GV xem chi tiết/danh sách lớp (GV-04) và xuất Excel (GV-12)."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.attendance.models import Attendance, AttendanceSession
from apps.auth.models import Role
from apps.grades.models import Grade
from apps.registration.models import Registration, TrangThaiDangKy
from apps.training.models import CurriculumCourse, TrangThaiLop
from core.factories import (
    ClassSectionFactory,
    CourseFactory,
    CurriculumFactory,
    SemesterFactory,
    StudentProfileFactory,
    UserFactory,
)

pytestmark = pytest.mark.django_db
BC = "/api/reports/training/"


def lop_voi_sv(n, hk, trang_thai=TrangThaiLop.MO, toi_da=60, **kw):
    lop = ClassSectionFactory(hoc_ky=hk, trang_thai=trang_thai, si_so_toi_da=toi_da, **kw)
    sv = [StudentProfileFactory() for _ in range(n)]
    for p in sv:
        Registration.objects.create(sinh_vien=p.user, lop_hoc_phan=lop)
    lop.si_so_hien_tai = n
    lop.save(update_fields=["si_so_hien_tai"])
    return lop, sv


def sheets(res):
    wb = load_workbook(BytesIO(res.content))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


@pytest.fixture
def pdt(login_as):
    return login_as(UserFactory(role=Role.PHONG_DAO_TAO))


# ------------------------------------------------------------------ báo cáo đào tạo


def test_bao_cao_dao_tao_dem_dung_va_sap_xep(pdt):
    hk = SemesterFactory()
    day, sv_chung = lop_voi_sv(2, hk, toi_da=2)  # đầy
    con_cho, _ = lop_voi_sv(3, hk, toi_da=10)  # còn chỗ
    lop_voi_sv(1, hk, TrangThaiLop.DANG_HOC, toi_da=1)  # không tính đầy / còn chỗ (đã hết đăng ký)
    ClassSectionFactory(hoc_ky=hk)  # KE_HOACH: bỏ qua
    # Một SV học thêm lớp khác → vẫn là 1 người trong "số sinh viên"
    Registration.objects.create(sinh_vien=sv_chung[0].user, lop_hoc_phan=con_cho)
    con_cho.si_so_hien_tai = 4
    con_cho.save(update_fields=["si_so_hien_tai"])
    lop_voi_sv(5, SemesterFactory())  # học kỳ khác: bỏ qua

    res = pdt.get(BC, {"hoc_ky": hk.id})
    assert res.status_code == 200
    d = res.data
    assert d["tong"] == {
        "so_sinh_vien": 2 + 3 + 1,
        "so_lop": 3,
        "tong_luot_dang_ky": 2 + 4 + 1,
        "lop_day": 1,
        "lop_con_cho": 1,
    }
    assert d["theo_lop"][-1]["ma_lop"] == con_cho.ma_lop  # lấp đầy thấp nhất xếp cuối
    assert d["theo_lop"][-1]["ty_le_lap_day"] == 40.0
    assert d["theo_mon"][0]["luot_dang_ky"] == 4


def test_bao_cao_dao_tao_loc_nganh_va_mon(pdt):
    hk = SemesterFactory()
    ct = CurriculumFactory()
    mon = CourseFactory()
    CurriculumCourse.objects.create(ctdt=ct, mon_hoc=mon, khoi_kien_thuc="CO_SO_NGANH")
    trong, _ = lop_voi_sv(2, hk, mon_hoc=mon)
    lop_voi_sv(3, hk)  # môn ngoài CTĐT của ngành
    res = pdt.get(BC, {"hoc_ky": hk.id, "nganh": ct.nganh_id})
    assert [r["ma_lop"] for r in res.data["theo_lop"]] == [trong.ma_lop]
    res = pdt.get(BC, {"hoc_ky": hk.id, "mon_hoc": mon.id})
    assert res.data["tong"]["so_lop"] == 1


def test_bao_cao_dao_tao_loi_tham_so_va_phan_quyen(login_as, pdt):
    assert pdt.get(BC, {"nganh": "abc"}).status_code == 400
    assert pdt.get(BC, {"hoc_ky": 999999}).status_code == 400
    for role in (Role.KE_TOAN, Role.SINH_VIEN, Role.GIAO_VIEN, Role.ADMIN):
        c = login_as(UserFactory(role=role))
        assert c.get(BC).status_code == 403
        assert c.get(BC + "export/").status_code == 403


def test_xuat_excel_bao_cao_dao_tao(pdt):
    hk = SemesterFactory()
    lop, _ = lop_voi_sv(2, hk, toi_da=4)
    res = pdt.get(BC + "export/", {"hoc_ky": hk.id})
    assert res.status_code == 200
    s = sheets(res)
    assert list(s) == ["Tổng quan", "Theo môn", "Theo lớp"]
    assert ["Tổng lượt đăng ký", 2] in s["Tổng quan"]
    assert s["Theo lớp"][1][0] == lop.ma_lop
    assert s["Theo lớp"][1][5:] == [50.0, "Mở"]


# ------------------------------------------------------------------ GV xem lớp mình


def gv_lop(login_as, **kw):
    lop = ClassSectionFactory(trang_thai=TrangThaiLop.DANG_HOC, **kw)
    return lop, login_as(lop.giao_vien.user)


def buoi(lop, ngay):
    return AttendanceSession.objects.create(
        lop_hoc_phan=lop, ngay_hoc=ngay, buoi="SANG", tiet_bat_dau=1
    )


def test_gv_xem_chi_tiet_lop_minh_va_bi_chan_lop_khac(login_as):
    lop, gv = gv_lop(login_as, mon_hoc=CourseFactory(tin_chi_ly_thuyet=2, tin_chi_thuc_hanh=1))
    buoi(lop, date(2026, 3, 2))
    res = gv.get(f"/api/class-sections/{lop.id}/")
    assert res.status_code == 200
    assert res.data["tong_so_buoi"] == 1
    assert (res.data["tin_chi_ly_thuyet"], res.data["tin_chi_thuc_hanh"]) == (2, 1)
    assert res.data["ngay_bat_dau_hieu_luc"] == date(2026, 2, 1)  # lớp chưa có ngày → theo học kỳ
    assert "han_nhap_diem_hieu_luc" in res.data

    khac = ClassSectionFactory()
    assert gv.get(f"/api/class-sections/{khac.id}/").status_code == 403
    assert gv.get(f"/api/class-sections/{khac.id}/students/").status_code == 403
    assert gv.get(f"/api/class-sections/{khac.id}/export/").status_code == 403
    assert gv.get("/api/class-sections/").status_code == 403  # GV không xem danh sách mọi lớp
    assert gv.patch(f"/api/class-sections/{lop.id}/", {"si_so_toi_da": 5}).status_code == 403


def test_pdt_kt_xem_chi_tiet_lop_con_sv_thi_khong(login_as, pdt):
    lop = ClassSectionFactory()
    url = f"/api/class-sections/{lop.id}/"
    assert pdt.get(url).status_code == 200
    assert login_as(UserFactory(role=Role.KE_TOAN)).get(url).status_code == 200
    assert login_as(UserFactory(role=Role.SINH_VIEN)).get(url).status_code == 403


def test_lop_cua_toi_kem_thoi_han_va_trang_thai_duyet_diem(login_as):
    lop, gv = gv_lop(login_as)
    sv = StudentProfileFactory()
    Registration.objects.create(sinh_vien=sv.user, lop_hoc_phan=lop)
    res = gv.get("/api/class-sections/my/")
    assert res.data["results"][0]["diem_da_duyet"] is False
    assert res.data["results"][0]["ten_hk"] == lop.hoc_ky.ten_hk
    assert res.data["results"][0]["han_nhap_diem_hieu_luc"] == date(2026, 7, 14)  # 30/06 + 14 ngày
    Grade.objects.create(lop_hoc_phan=lop, sinh_vien=sv.user, trang_thai="KHOA")
    assert gv.get("/api/class-sections/my/").data["results"][0]["diem_da_duyet"] is True


def test_gv_danh_sach_sv_chi_gom_dang_dang_ky(login_as):
    lop, gv = gv_lop(login_as, si_so_toi_da=40)
    a, b = StudentProfileFactory(ma_sv="SV001"), StudentProfileFactory(ma_sv="SV002")
    Registration.objects.create(sinh_vien=b.user, lop_hoc_phan=lop)
    Registration.objects.create(sinh_vien=a.user, lop_hoc_phan=lop)
    da_huy = StudentProfileFactory()
    Registration.objects.create(
        sinh_vien=da_huy.user, lop_hoc_phan=lop, trang_thai=TrangThaiDangKy.DA_HUY
    )
    res = gv.get(f"/api/class-sections/{lop.id}/students/")
    assert res.status_code == 200
    assert (res.data["si_so"], res.data["si_so_toi_da"]) == (2, 40)
    assert [r["ma_sv"] for r in res.data["records"]] == ["SV001", "SV002"]  # theo mã SV
    assert res.data["records"][0]["nganh_ten"] == a.nganh.ten_nganh


def test_gv_xuat_excel_lop_ba_sheet(login_as):
    lop, gv = gv_lop(login_as)
    a, b = StudentProfileFactory(ma_sv="SV001"), StudentProfileFactory(ma_sv="SV002")
    for p in (a, b):
        Registration.objects.create(sinh_vien=p.user, lop_hoc_phan=lop)
    Grade.objects.create(
        lop_hoc_phan=lop,
        sinh_vien=a.user,
        diem_qua_trinh=8,
        diem_giua_ky=7,
        diem_cuoi_ky=9,
        diem_tong_ket=8.2,
        diem_chu="B",
        trang_thai="CHINH_THUC",
    )
    for b_dd, tt, tre in (
        (buoi(lop, date(2026, 3, 2)), "CO_MAT", True),
        (buoi(lop, date(2026, 3, 9)), "VANG_KHONG_PHEP", False),
    ):
        Attendance.objects.create(
            attendance_session=b_dd,
            lop_hoc_phan=lop,
            ngay_hoc=b_dd.ngay_hoc,
            sinh_vien=a.user,
            trang_thai=tt,
            di_tre=tre,
        )

    res = gv.get(f"/api/class-sections/{lop.id}/export/")
    assert res.status_code == 200
    assert lop.ma_lop in res["Content-Disposition"]
    s = sheets(res)
    assert list(s) == ["Danh sách SV", "Bảng điểm", "Điểm danh"]
    assert [r[1] for r in s["Danh sách SV"][1:]] == ["SV001", "SV002"]
    assert s["Bảng điểm"][1][2:] == [8.0, 7.0, 9.0, 8.2, "B", "Chính thức"]
    assert s["Bảng điểm"][2][-1] == "Chưa nhập"
    dd = s["Điểm danh"]
    assert dd[0][2:4] == ["02/03 T1", "09/03 T1"]
    assert dd[1][2:] == ["T", "V", 1, 0, 50.0]  # đi trễ · vắng KP · 1/2 buổi = 50%
    assert dd[2][2:4] == [None, None]  # SV chưa có bản ghi điểm danh


def test_pdt_xuat_excel_lop_bat_ky_con_ke_toan_sinh_vien_thi_khong(login_as, pdt):
    lop = ClassSectionFactory(trang_thai=TrangThaiLop.DANG_HOC)
    url = f"/api/class-sections/{lop.id}/export/"
    res = pdt.get(url)
    assert res.status_code == 200 and lop.ma_lop in res["Content-Disposition"]
    assert pdt.get(f"/api/class-sections/{lop.id}/students/").status_code == 403  # GV-04 chỉ của GV
    for vai_tro in (Role.KE_TOAN, Role.SINH_VIEN, Role.ADMIN):
        assert login_as(UserFactory(role=vai_tro)).get(url).status_code == 403
