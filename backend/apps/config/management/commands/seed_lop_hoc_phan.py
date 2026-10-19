"""Nạp LỚP HỌC PHẦN + lịch học cho mọi học kỳ đã có, theo các khối "Học kỳ N - Năm học ..." trong file Excel CTĐT.

Mỗi môn mở 4–7 lớp trong kỳ (nhiều CTĐT cùng học môn thì nhiều lớp hơn), mỗi lớp một giảng viên khác nhau
(của khoa phụ trách môn) và một phòng.
Ràng buộc đảm bảo (và được kiểm lại sau khi nạp, sai thì rollback):
  - một giảng viên không dạy hai lớp trùng giờ; một phòng không có hai lớp trùng giờ;
  - với mỗi (CTĐT, học kỳ), các lớp L01 của mọi môn đôi một không trùng giờ, L02 cũng vậy, L03 cũng vậy
    → luôn có ít nhất SO_TKB thời khóa biểu đăng ký đủ mọi môn của kỳ mà không trùng lịch.
Chạy lại an toàn khi chưa có đăng ký (kết quả cố định); đã có đăng ký thì từ chối để không đổi lịch của sinh viên.
"""

import re
from collections import Counter, defaultdict
from datetime import timedelta

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.config.management.commands.seed_ctdt import NGANH, THU_MUC
from apps.profiles.models import TeacherProfile
from apps.registration.models import Registration
from apps.training.models import (
    ClassSection,
    Course,
    Schedule,
    Semester,
    TrangThaiHocKy,
    TrangThaiLop,
)

SO_LOP_TOI_THIEU, SO_LOP_TOI_DA = 4, 7
SO_TKB = 3  # số "lớp thứ j" đầu tiên được đảm bảo thành một thời khóa biểu đầy đủ không trùng
TOI_DA_TIET_TUAN = 8  # môn thực tập/đồ án trong file có tới 20 tiết/tuần — gói lại cho vừa một tuần học
BUOI = [(1, 5), (6, 10), (11, 15)]  # sáng, chiều, tối (BR-08)
THU = [2, 3, 4, 5, 6, 7]
THU_TU_BUOI = [(0, 1, 2), (1, 0, 2), (0, 1, 2)]  # theo lớp: lớp 2 ưu tiên buổi chiều
PHONG_LY_THUYET = [f"{toa}{tang}{p:02d}" for toa in "AB" for tang in range(1, 5) for p in range(1, 11)]
PHONG_THUC_HANH = [f"LAB{n:02d}" for n in range(1, 31)]
TEN_HK = re.compile(r"Học kỳ (\d) - Năm học (\d{4}) - (\d{4})")


def doc_nhom():
    """{ten_hk: {ma_nganh: [ma_mon]}} theo thứ tự khối "Học kỳ N - Năm học ..." trong file Excel."""
    nhom = defaultdict(dict)
    for file, ma_nganh, *_ in NGANH:
        ws = openpyxl.load_workbook(THU_MUC / file, data_only=True).active
        dau = next(ws.iter_rows(values_only=True))
        cot = {ten: i for i, ten in enumerate(dau) if ten}
        hk = None
        for r in ws.iter_rows(min_row=2, values_only=True):
            ma, ten = r[cot["Mã MH"]], r[cot["Tên môn học"]]
            if not ma:
                if ten:
                    hk = ten.strip()
                continue
            nhom[hk].setdefault(ma_nganh, []).append(ma)
    return nhom


def cac_buoi_hoc(mon):
    """Số tiết của từng buổi trong tuần: tổng tiết / 15 tuần, tối đa TOI_DA_TIET_TUAN, tách 2 buổi nếu > 5 tiết."""
    tuan = -(-(mon.so_tiet_ly_thuyet + mon.so_tiet_thuc_hanh) // 15)
    w = max(1, min(TOI_DA_TIET_TUAN, tuan))
    return [w] if w <= 5 else [w - w // 2, w // 2]


def giao(a, b):
    """Hai buổi (thu, tiết đầu, tiết cuối) trùng nhau — cùng định nghĩa BR-03 (registration.services.lich_trung)."""
    return a[0] == b[0] and a[1] <= b[2] and b[1] <= a[2]


def ung_vien(do_dai, lop, ma):
    """Các (thu, tiết đầu, tiết cuối) có thể đặt một buổi `do_dai` tiết, theo thứ tự ưu tiên của lớp."""
    xoay = (lop * 2 + sum(map(ord, ma))) % len(THU)
    ngay = THU[xoay:] + THU[:xoay]
    for b in THU_TU_BUOI[lop % len(THU_TU_BUOI)]:
        dau, cuoi = BUOI[b]
        for thu in ngay:
            for s in range(dau, cuoi - do_dai + 2):
                yield (thu, s, s + do_dai - 1)


class KeHoach:
    """Trạng thái xếp lịch của MỘT học kỳ."""

    def __init__(self, nhom_hk, mon, gv_theo_khoa):
        self.mon = mon
        self.gv_theo_khoa = gv_theo_khoa
        self.xung_dot = defaultdict(set)  # môn → các môn cùng CTĐT trong kỳ (không được trùng giờ cùng "lớp thứ j")
        so_ctdt = Counter()  # môn → số CTĐT cùng học môn này trong kỳ
        for ds in nhom_hk.values():
            for ma in ds:
                self.xung_dot[ma] |= set(ds) - {ma}
                so_ctdt[ma] += 1
        # 4–7 lớp/môn: nhiều CTĐT chung môn thì nhiều lớp hơn (+0/1 cố định theo mã để các môn không đồng loạt bằng nhau).
        self.so_lop = {
            ma: min(SO_LOP_TOI_DA, SO_LOP_TOI_THIEU - 1 + so_ctdt[ma] + sum(map(ord, ma)) % 2)
            for ma in self.xung_dot
        }
        self.giu = defaultdict(dict)  # lớp thứ j → {ma_mon: [buổi]}
        self.gv_ban = defaultdict(list)
        self.phong_ban = defaultdict(list)
        self.tai = Counter()
        self.gv_cua_mon = defaultdict(set)
        self.ket_qua = {}  # (ma_mon, j) → (gv, phong, [buổi])

    def ranh_gv(self, gv, buoi):
        return not any(giao(b, x) for b in buoi for x in self.gv_ban[gv])

    def ranh_phong(self, phong, buoi):
        return not any(giao(b, x) for b in buoi for x in self.phong_ban[phong])

    def hop_le_nhom(self, j, ma, b):
        return not any(giao(b, x) for d in self.xung_dot[ma] for x in self.giu[j].get(d, []))

    def dat(self, ma, j):
        mon = self.mon[ma]
        do_dai = cac_buoi_hoc(mon)
        phong_ds = PHONG_THUC_HANH if mon.so_tiet_ly_thuyet == 0 else PHONG_LY_THUYET
        gv_ds = sorted(self.gv_theo_khoa[mon.khoa_id], key=lambda g: (self.tai[g.id], g.id))
        gv_ds = [g for g in gv_ds if g.id not in self.gv_cua_mon[ma]]
        for b1 in ung_vien(do_dai[0], j, ma):
            if not self.hop_le_nhom(j, ma, b1):
                continue
            kha_di = [[b1]]
            if len(do_dai) == 2:
                kha_di = [
                    [b1, b2]
                    for b2 in ung_vien(do_dai[1], j, ma)
                    if b2[0] != b1[0] and self.hop_le_nhom(j, ma, b2)
                ]
            for buoi in kha_di:
                gv = next((g for g in gv_ds if self.ranh_gv(g.id, buoi)), None)
                phong = next((p for p in phong_ds if self.ranh_phong(p, buoi)), None)
                if gv and phong:
                    self.giu[j].setdefault(ma, []).extend(buoi)
                    self.gv_ban[gv.id].extend(buoi)
                    self.phong_ban[phong].extend(buoi)
                    self.tai[gv.id] += 1
                    self.gv_cua_mon[ma].add(gv.id)
                    self.ket_qua[(ma, j)] = (gv, phong, buoi)
                    return
        raise CommandError(f"Không xếp được lịch cho {ma} (lớp thứ {j + 1}).")

    def xep(self):
        thu_tu = sorted(self.xung_dot.keys(), key=lambda m: (-sum(cac_buoi_hoc(self.mon[m])), m))
        for j in range(SO_LOP_TOI_DA):
            for ma in thu_tu:
                if j < self.so_lop[ma]:
                    self.dat(ma, j)
        return self.ket_qua


class Command(BaseCommand):
    help = "Nạp lớp học phần + lịch cho mọi học kỳ đã có (4–7 lớp/môn, GV và phòng không trùng giờ)."

    @transaction.atomic
    def handle(self, *args, **options):
        if Registration.objects.exists():
            raise CommandError("Đã có đăng ký — không nạp lại lớp học phần để khỏi đổi lịch của sinh viên.")
        nhom = doc_nhom()
        mon = {c.ma_mon: c for c in Course.objects.all()}
        gv_theo_khoa = defaultdict(list)
        for g in TeacherProfile.objects.order_by("id"):
            gv_theo_khoa[g.khoa_id].append(g)

        tong_lop = 0
        for hk in Semester.objects.order_by("ngay_bat_dau"):
            if hk.ten_hk not in nhom:
                continue
            ket_qua = KeHoach(nhom[hk.ten_hk], mon, gv_theo_khoa).xep()
            ky, nam1, nam2 = TEN_HK.fullmatch(hk.ten_hk).groups()
            trang_thai = TrangThaiLop.MO if hk.trang_thai == TrangThaiHocKy.MO else TrangThaiLop.DA_DONG
            ma_lop_moi = set()
            for (ma, j), (gv, phong, buoi) in sorted(ket_qua.items()):
                m = mon[ma]
                ma_lop = f"{ma}.L{j + 1:02d}-{nam1[2:]}{nam2[2:]}.{ky}"
                ma_lop_moi.add(ma_lop)
                lop, _ = ClassSection.objects.update_or_create(
                    ma_lop=ma_lop,
                    defaults={
                        "mon_hoc": m,
                        "hoc_ky": hk,
                        "giao_vien": gv,
                        "si_so_toi_da": 40 if m.so_tiet_ly_thuyet == 0 else 60,
                        "ngay_bat_dau": hk.ngay_bat_dau,
                        "ngay_ket_thuc": hk.ngay_ket_thuc,
                        "han_nhap_diem": hk.ngay_ket_thuc + timedelta(days=14),
                        "trang_thai": trang_thai,
                    },
                )
                lop.lich_hoc.all().delete()
                Schedule.objects.bulk_create(
                    Schedule(
                        lop_hoc_phan=lop,
                        thu=thu,
                        tiet_bat_dau=dau,
                        so_tiet_moi_buoi=cuoi - dau + 1,
                        phong=phong,
                    )
                    for thu, dau, cuoi in buoi
                )
            # Bỏ lớp cũ không còn trong kế hoạch (chạy lại sau khi đổi số lớp/môn); đã chặn khi có đăng ký.
            ClassSection.objects.filter(hoc_ky=hk).exclude(ma_lop__in=ma_lop_moi).delete()
            tong_lop += len(ket_qua)
            so_mon = len({ma for ma, _ in ket_qua})
            self.stdout.write(f"{hk.ten_hk}: {len(ket_qua)} lớp ({so_mon} môn)")
        self._kiem_tra(nhom)
        self.stdout.write(f"Tổng {tong_lop} lớp học phần. Kiểm tra trùng lịch GV/phòng và thời khóa biểu: ĐẠT.")

    def _kiem_tra(self, nhom):
        """Dò lại từ DB (không dựa vào bộ nhớ lúc xếp): lỗi thì CommandError → rollback."""
        theo_hk = defaultdict(list)  # ten_hk → [(lớp, [lịch])]
        for lop in ClassSection.objects.select_related("hoc_ky").prefetch_related("lich_hoc"):
            theo_hk[lop.hoc_ky.ten_hk].append((lop, list(lop.lich_hoc.all())))

        def buoi(s):
            return (s.thu, s.tiet_bat_dau, s.tiet_ket_thuc)

        for hk, ds in theo_hk.items():
            for ten, khoa_cua in (("giảng viên", lambda lop, s: lop.giao_vien_id), ("phòng", lambda lop, s: s.phong)):
                theo = defaultdict(list)
                for lop, lich in ds:
                    for s in lich:
                        theo[khoa_cua(lop, s)].append((lop.ma_lop, buoi(s)))
                for khoa, cac in theo.items():
                    for i, (ma_a, a) in enumerate(cac):
                        for ma_b, b in cac[i + 1 :]:
                            if ma_a != ma_b and giao(a, b):
                                raise CommandError(f"{hk}: {ten} {khoa} trùng giờ giữa {ma_a} và {ma_b}.")

        for hk, theo_nganh in nhom.items():
            if hk not in theo_hk:
                continue
            lich_lop = {lop.ma_lop: [buoi(s) for s in lich] for lop, lich in theo_hk[hk]}
            so_lop = Counter(ma_lop.split(".L")[0] for ma_lop in lich_lop)
            for ma, so in so_lop.items():
                if not SO_LOP_TOI_THIEU <= so <= SO_LOP_TOI_DA:
                    raise CommandError(f"{hk}: môn {ma} có {so} lớp, ngoài khoảng {SO_LOP_TOI_THIEU}–{SO_LOP_TOI_DA}.")
            for ma_nganh, ds in theo_nganh.items():
                hop_le = 0
                for j in range(SO_TKB):
                    lop_j = [
                        ma_lop
                        for ma_lop in lich_lop
                        if f".L{j + 1:02d}-" in ma_lop and ma_lop.split(".L")[0] in ds
                    ]
                    if len(lop_j) != len(ds):
                        raise CommandError(f"{hk}/{ma_nganh}: thiếu lớp thứ {j + 1} của một số môn.")
                    tat_ca = [b for m in lop_j for b in lich_lop[m]]
                    if not any(giao(a, b) for i, a in enumerate(tat_ca) for b in tat_ca[i + 1 :]):
                        hop_le += 1
                if hop_le < SO_TKB:
                    raise CommandError(f"{hk}/{ma_nganh}: chỉ có {hop_le}/{SO_TKB} thời khóa biểu không trùng.")
