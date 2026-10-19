"""Nạp 5 ngành THẬT + CTĐT (D23, riêng QTKD là D24) từ data/curriculums/*.xlsx. Chạy lại an toàn.

Tách khỏi seed_demo. Nếu chạy lại seed_demo thì chạy lệnh này SAU nó: ngành KTPM
dùng chung mã với ngành demo nên seed_demo sẽ đặt lại khoa/số năm của KTPM.
"""

import re
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.core.management.base import BaseCommand
from django.db import transaction

from apps.training.models import (
    Course,
    Curriculum,
    CurriculumCourse,
    Faculty,
    KhoiKienThuc,
    LoaiMon,
    Major,
    Prerequisite,
)

# Môn điều kiện (GDQP-AN, GDTC): bắt buộc để tốt nghiệp nhưng chấm Đạt/Không đạt, không tính GPA.
MON_DIEU_KIEN = {"KTCH002", "KTCH003", "KTCH004", "KTCH007", "KTCH013"}
THU_MUC = Path(__file__).resolve().parents[4] / "data" / "curriculums"

# (file, mã ngành, tên ngành, mã khoa/viện, số năm, tổng tín chỉ yêu cầu, khóa tuyển sinh)
NGANH = [
    ("kinhte.xlsx", "QTKD", "Quản trị kinh doanh", "KT", 4, 120, "D24"),
    ("ktpm.xlsx", "KTPM", "Kỹ thuật phần mềm", "VCNS", 5, 150, "D23"),
    ("oto.xlsx", "OTO", "Kỹ thuật ô tô", "VKTCN", 4, 120, "D23"),
    ("supham.xlsx", "GDTH", "Giáo dục Tiểu học", "SP", 4, 120, "D23"),
    ("tnmt.xlsx", "QLMT", "Quản lý tài nguyên môi trường", "VCNXBV", 4, 120, "D23"),
]

# (môn, môn tiên quyết) — suy luận theo thứ tự học chuẩn của ngành; PĐT chỉnh lại trên màn Tiên quyết.
# Mỗi cặp đã kiểm: môn tiên quyết có trong mọi CTĐT chứa môn đó và nằm ở học kỳ gợi ý sớm hơn.
TIEN_QUYET = [
    # KTPM (PĐT chỉ định)
    ("LING105", "LING022"),  # Kỹ thuật lập trình ← Cơ sở lập trình
    ("LING010", "LING105"),  # Cấu trúc dữ liệu và giải thuật ← Kỹ thuật lập trình
    ("LING345", "LING344"),  # Toán cao cấp A2 ← A1
    ("LING396", "LING345"),  # Xác suất thống kê ← Toán cao cấp A2
    # QTKD
    ("LING096", "LING095"),  # Kinh tế vĩ mô ← vi mô
    ("QTKD019", "LING095"),  # Kinh tế học quản lý ← vi mô
    ("MKTG005", "LING166"),  # Kế toán tài chính ← Nguyên lý kế toán
    ("QTKD016", "MKTG005"),  # Phân tích BCTC ← Kế toán tài chính
    ("LING462", "LING238"),  # Tài chính doanh nghiệp ← Tài chính tiền tệ
    ("LING216", "LING456"),  # Quản trị chiến lược ← Quản trị học
    ("TCNH034", "LING169"),  # Kinh tế lượng ← Nguyên lý thống kê kinh tế
    ("QTKD003", "LING138"),  # Digital Marketing ← Marketing căn bản
    ("QTKD011", "LING138"),  # Quản trị thương hiệu ← Marketing căn bản
    ("LING152", "LING138"),  # Nghiên cứu thị trường ← Marketing căn bản
    ("QTKD002", "LING138"),  # Hành vi khách hàng ← Marketing căn bản
    # OTO
    ("LING017", "LING387"),  # Cơ kỹ thuật ← Vật lý đại cương A1
    ("LING235", "LING017"),  # Sức bền vật liệu ← Cơ kỹ thuật
    ("LING164", "LING235"),  # Nguyên lý chi tiết máy ← Sức bền vật liệu
    ("LING254", "LING074"),  # Thiết kế mô phỏng trên máy tính ← Hình họa vẽ kỹ thuật
    ("CNOT004", "LING055"),  # Cơ sở CN chế tạo máy ← Dung sai và kỹ thuật đo
    ("CNOT015", "LING041"),  # HT điện - điện tử trên ô tô ← Điện - điện tử cơ bản
    ("CNOT035", "LING041"),  # Vi điều khiển ứng dụng ← Điện - điện tử cơ bản
    ("CNOT011", "CNOT012"),  # Đồ án chuyên ngành ← Đồ án môn học
    # GDTH
    ("LING240", "LING239"),  # Tâm lý học lứa tuổi và SP ← Tâm lý học
    ("GDTH028", "GDTH007"),  # LL dạy học Toán ← Cơ sở Toán học 2
    ("GDTH027", "GDTH005"),  # LL dạy học Tiếng Việt ← Cơ sở Tiếng Việt 2
    ("GDTH025", "GDTH009"),  # LL dạy học TNXH ← Cơ sở TNXH 2
    # QLMT
    ("QLMT035", "QLMT013"),  # Phân tích môi trường ← Hóa phân tích
    ("LING213", "QLMT035"),  # Quan trắc môi trường ← Phân tích môi trường
    ("LING039", "QLMT019"),  # Đánh giá tác động MT ← Luật và chính sách MT
    ("LING029", "LING021"),  # Cơ sở viễn thám ← Cơ sở GIS
    ("LING092", "LING034"),  # Kiểm soát ô nhiễm MT ← Đại cương công nghệ MT
    # --- Suy luận sâu hơn (cùng cách kiểm: thứ tự học kỳ + có mặt trong CTĐT) ---
    # KTPM
    ("LING068", "LING020"),  # Hệ QTCSDL ← Cơ sở dữ liệu
    ("LING109", "LING256"),  # Lập trình Web ← Thiết kế Web
    ("LING196", "LING105"),  # PP lập trình HĐT ← Kỹ thuật lập trình
    ("LING110", "LING105"),  # Lập trình trên Windows ← Kỹ thuật lập trình
    ("LING135", "LING349"),  # Lý thuyết đồ thị ← Toán rời rạc
    ("KTPM031", "LING196"),  # Phân tích, thiết kế HĐT ← PP lập trình HĐT
    ("KTPM034", "LING031"),  # Kiến trúc và thiết kế PM ← Công nghệ phần mềm
    ("LING189", "LING196"),  # PT ứng dụng di động ← PP lập trình HĐT
    ("LING190", "LING189"),  # PT ứng dụng di động đa nền tảng ← PT ứng dụng di động
    ("KTPM023", "LING189"),  # Lập trình game trên di động ← PT ứng dụng di động
    ("LING358", "LING010"),  # Trí tuệ nhân tạo ← CTDL và giải thuật
    ("LING081", "LING358"),  # Học máy ← Trí tuệ nhân tạo
    ("KTPM024", "LING137"),  # Lập trình mạng ← Mạng máy tính
    ("CNTT034", "LING005"),  # An toàn ứng dụng web ← An toàn và bảo mật thông tin
    ("CNTT043", "LING093"),  # Hệ điều hành ← Kiến trúc máy tính
    ("LING042", "LING137"),  # Điện toán đám mây ← Mạng máy tính
    ("KTPM012", "LING196"),  # Công nghệ Java ← PP lập trình HĐT
    ("KTPM011", "LING196"),  # Công nghệ .NET ← PP lập trình HĐT
    ("KTPM015", "KTPM033"),  # Đồ án chuyên ngành ← Đồ án cơ sở ngành
    ("LING210", "LING031"),  # Quản lý dự án CNTT ← Công nghệ phần mềm
    ("KTPM030", "LING109"),  # PT phần mềm hướng dịch vụ ← Lập trình Web
    # QTKD
    ("LING225", "LING456"),  # Quản trị nguồn nhân lực ← Quản trị học
    ("LING400", "LING456"),  # Nghệ thuật lãnh đạo ← Quản trị học
    ("LING226", "LING456"),  # Quản trị sản xuất ← Quản trị học
    ("LING220", "LING456"),  # Quản trị dự án ← Quản trị học
    ("LING215", "LING456"),  # Quản trị chất lượng ← Quản trị học
    ("LING463", "LING226"),  # Quản trị chuỗi cung ứng ← Quản trị sản xuất
    ("QTKD025", "LING226"),  # Sản xuất tinh gọn ← Quản trị sản xuất
    ("QTKD026", "LING226"),  # QT sản xuất và điều hành ← Quản trị sản xuất
    ("TMDT021", "LING238"),  # Quản trị tài chính ← Tài chính tiền tệ
    ("QTKD017", "LING238"),  # Chính sách tài chính tiền tệ ← Tài chính tiền tệ
    ("LING440", "LING096"),  # Kinh tế phát triển ← Kinh tế vĩ mô
    ("LING182", "LING166"),  # Phân tích HĐKD ← Nguyên lý kế toán
    ("QTKD021", "MKTG005"),  # BCTC và phân tích BCTC ← Kế toán tài chính
    ("QTKD022", "TCNH034"),  # Phân tích và dự báo KD ← Kinh tế lượng
    ("QTKD012", "LING247"),  # QT xuất nhập khẩu ← Thanh toán quốc tế
    ("LING454", "LING138"),  # Quản trị bán hàng ← Marketing căn bản
    # OTO
    ("CNOT024", "LING017"),  # Lý thuyết ô tô ← Cơ kỹ thuật
    ("CNOT021", "CNOT037"),  # Tính toán ô tô ← Kết cấu ô tô
    ("CNOT006", "CNOT037"),  # CN chẩn đoán và sửa chữa ô tô ← Kết cấu ô tô
    ("CNOT016", "CNOT035"),  # HT điều khiển tự động ô tô ← Vi điều khiển ứng dụng
    ("CNOT017", "CNOT016"),  # HT điều khiển và giám sát ô tô ← HT điều khiển tự động ô tô
    ("CNOT020", "CNOT006"),  # Kiểm định và chẩn đoán ô tô ← CN chẩn đoán và sửa chữa
    ("CNOT001", "LING254"),  # Ứng dụng máy tính thiết kế ô tô ← Thiết kế mô phỏng
    # GDTH
    ("GDTH045", "GDTH007"),  # TCHĐ dạy học Toán ← Cơ sở Toán học 2
    ("GDTH046", "GDTH005"),  # TCHĐ dạy học Tiếng Việt ← Cơ sở Tiếng Việt 2
    ("GDTH051", "GDTH009"),  # TCHĐ dạy học TNXH ← Cơ sở TNXH 2
    ("GDTH013", "GDTH025"),  # Dạy học TNXH tích hợp ← LL dạy học TNXH
    ("GDTH064", "GDTH028"),  # Bồi dưỡng HS giỏi Toán ← LL dạy học Toán
    ("GDTH065", "GDTH028"),  # Phát triển tư duy qua dạy Toán ← LL dạy học Toán
    # QLMT
    ("QLMT012", "QLMT013"),  # Hóa môi trường ← Hóa phân tích
    ("LING141", "LING452"),  # Mô hình hóa môi trường ← Các phương pháp thống kê
    ("LING229", "LING086"),  # Quy hoạch sử dụng đất ← Khoa học đất
    ("QLMT027", "LING092"),  # QL MT không khí và tiếng ồn ← Kiểm soát ô nhiễm
]
# Môn cùng tên, hậu tố số N ≥ 2 (Toán cao cấp C2, Thực tập sư phạm 3…) ← môn hậu tố N-1 trong cùng CTĐT,
# chỉ khi môn N-1 nằm ở học kỳ gợi ý sớm hơn (cùng học kỳ = học song song, không phải tiên quyết).
HAU_TO_SO = re.compile(r"^(.*\D)(\d+)$")

# Khối kiến thức: suy ra theo quy tắc (PĐT chỉnh lại trên màn CTĐT nếu sai).
CHUNG_DAI_CUONG = {"LING344", "LING345", "LING346", "LING347", "LING387", "LING396"}  # Toán, Lý, XSTK
TIEN_TO_CHUYEN_NGANH = {"QTKD": "QTKD", "KTPM": "KTPM", "OTO": "CNOT", "GDTH": "GDTH", "QLMT": "QLMT"}

TIN_CHI = re.compile(r"\s*\((\d+)\+(\d+)\)\s*$")  # "Tên môn (2+1)" → lý thuyết 2, thực hành 1


class Command(BaseCommand):
    help = "Nạp 5 ngành thật + CTĐT từ file Excel."

    @transaction.atomic
    def handle(self, *args, **options):
        call_command("seed_khoa_vien", verbosity=0)
        for file, ma_nganh, ten_nganh, ma_khoa, so_nam, tong_tc, khoa_ts in NGANH:
            khoa = Faculty.objects.get(ma_khoa=ma_khoa)
            nganh, _ = Major.objects.update_or_create(
                ma_nganh=ma_nganh,
                defaults={"ten_nganh": ten_nganh, "thoi_gian_dao_tao": so_nam, "khoa": khoa},
            )
            ctdt, _ = Curriculum.objects.update_or_create(
                ma_ctdt=f"CT-{ma_nganh}-{khoa_ts}",
                defaults={
                    "nganh": nganh,
                    "khoa_tuyen_sinh": khoa_ts,
                    "tong_tin_chi_yeu_cau": tong_tc,
                },
            )
            so_mon = self._nap_mon(THU_MUC / file, ctdt, khoa)
            self._phan_khoi(ctdt, TIEN_TO_CHUYEN_NGANH[ma_nganh])
            self.stdout.write(f"{ma_nganh}: {so_mon} môn ({khoa.ten_khoa}, {so_nam} năm, {tong_tc} TC)")
        self._nap_tien_quyet()

    def _nap_tien_quyet(self):
        cap = set(TIEN_QUYET)
        for ctdt in Curriculum.objects.all():
            mon = {
                cc.mon_hoc.ten_mon: (cc.mon_hoc.ma_mon, cc.hoc_ky_goi_y)
                for cc in ctdt.mon_ctdt.select_related("mon_hoc")
            }
            for ten, (ma, hk) in mon.items():
                m = HAU_TO_SO.match(ten)
                truoc = mon.get(f"{m[1]}{int(m[2]) - 1}") if m and int(m[2]) >= 2 else None
                if truoc and truoc[1] < hk:
                    cap.add((ma, truoc[0]))
        for ma, ma_tq in sorted(cap):
            Prerequisite.objects.get_or_create(
                mon_hoc=Course.objects.get(ma_mon=ma), mon_tien_quyet=Course.objects.get(ma_mon=ma_tq)
            )
        self.stdout.write(f"Môn tiên quyết: {len(cap)} cặp.")

    def _phan_khoi(self, ctdt, tien_to):
        """Tốt nghiệp > Chung (KTCH*, Toán/Lý/XSTK) > Tự chọn > Chuyên ngành (mã cùng tiền tố ngành) > Cơ sở ngành.
        "Thực hành X" theo khối của X."""

        def khoi_cua(cc):
            ma, ten = cc.mon_hoc.ma_mon, cc.mon_hoc.ten_mon
            if "tốt nghiệp" in ten.lower():
                return KhoiKienThuc.TOT_NGHIEP
            if ma.startswith("KTCH") or ma in CHUNG_DAI_CUONG:
                return KhoiKienThuc.CHUNG
            if cc.loai_mon == LoaiMon.TU_CHON:
                return KhoiKienThuc.TU_CHON
            return KhoiKienThuc.CHUYEN_NGANH if ma.startswith(tien_to) else KhoiKienThuc.CO_SO_NGANH

        ccs = list(ctdt.mon_ctdt.select_related("mon_hoc"))
        khoi_theo_ten = {cc.mon_hoc.ten_mon: khoi_cua(cc) for cc in ccs}
        for cc in ccs:
            ten = cc.mon_hoc.ten_mon
            khoi = khoi_theo_ten.get(ten.removeprefix("Thực hành "), khoi_cua(cc))
            if cc.khoi_kien_thuc != khoi:
                cc.khoi_kien_thuc = khoi
                cc.save(update_fields=["khoi_kien_thuc"])

    def _nap_mon(self, duong_dan, ctdt, khoa):
        ws = openpyxl.load_workbook(duong_dan, data_only=True).active
        dau = next(ws.iter_rows(values_only=True))
        cot = {ten: i for i, ten in enumerate(dau) if ten}
        hoc_ky = 0
        dem = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            ma, ten = r[cot["Mã MH"]], r[cot["Tên môn học"]]
            if not ma:
                if ten:  # dòng "Học kỳ N - Năm học ..." mở học kỳ gợi ý kế tiếp
                    hoc_ky += 1
                continue
            tin_chi = sum(map(int, TIN_CHI.search(ten).groups()))
            # Chỉ môn có chữ "Thực hành" trong tên mới tính tín chỉ thực hành (giá TH); còn lại là lý thuyết.
            lt, th = (0, tin_chi) if "thực hành" in ten.lower() else (tin_chi, 0)
            # Môn dùng chung giữa các ngành: giữ khoa của ngành nạp đầu, không giành lại.
            mon, _ = Course.objects.get_or_create(
                ma_mon=ma,
                defaults={
                    "ten_mon": TIN_CHI.sub("", ten),
                    "tong_tin_chi": tin_chi,
                    "tin_chi_ly_thuyet": lt,
                    "tin_chi_thuc_hanh": th,
                    "so_tiet_ly_thuyet": r[cot["lý thuyết"]] or 0,
                    "so_tiet_thuc_hanh": r[cot["thực hành"]] or 0,
                    "khoa": khoa,
                    "loai_danh_gia": "DAT_KHONG_DAT" if ma in MON_DIEU_KIEN else "DIEM",
                },
            )
            if (mon.tin_chi_ly_thuyet, mon.tin_chi_thuc_hanh) != (lt, th):  # chạy lại sau khi đổi quy tắc
                mon.tin_chi_ly_thuyet, mon.tin_chi_thuc_hanh = lt, th
                mon.save(update_fields=["tin_chi_ly_thuyet", "tin_chi_thuc_hanh"])
            bat_buoc = r[cot["Môn bắt buộc"]] == "x"
            CurriculumCourse.objects.update_or_create(
                ctdt=ctdt,
                mon_hoc=mon,
                defaults={
                    "loai_mon": LoaiMon.BAT_BUOC if bat_buoc else LoaiMon.TU_CHON,
                    "hoc_ky_goi_y": hoc_ky,
                },
            )
            dem += 1
        return dem
