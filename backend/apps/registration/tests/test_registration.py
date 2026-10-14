"""Đăng ký / hủy môn (BR-01..07, BR-05, BR-06), idempotency, học phí sau đăng ký, thông báo."""

import uuid
from datetime import timedelta
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.audit.models import AuditLog
from apps.auth.models import Role
from apps.grades.models import Grade, TrangThaiDiem
from apps.notifications.models import Notification
from apps.registration.models import Registration, TrangThaiDangKy
from apps.training.models import (
    ClassSection,
    CurriculumCourse,
    Prerequisite,
    TrangThaiLop,
)
from apps.tuition.models import HocPhanGia, Tuition
from core.factories import (
    ClassSectionFactory,
    CourseFactory,
    CurriculumFactory,
    ScheduleFactory,
    SemesterFactory,
    StudentProfileFactory,
    UserFactory,
)

pytestmark = pytest.mark.django_db
DK = "/api/registrations/"


def khoa_idem():
    return {"HTTP_IDEMPOTENCY_KEY": str(uuid.uuid4())}


@pytest.fixture
def hk():
    now = timezone.now()
    return SemesterFactory(
        cua_so_dk_bat_dau=now - timedelta(days=1), cua_so_dk_ket_thuc=now + timedelta(days=7)
    )


@pytest.fixture
def sv():
    ctdt = CurriculumFactory()
    return StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt)


@pytest.fixture
def sv_api(login_as, sv):
    return login_as(sv.user)


def tao_lop(hk, mon=None, thu=2, tiet=1, so_tiet=3, **kw):
    kw.setdefault("trang_thai", TrangThaiLop.MO)
    if mon is not None:
        kw["mon_hoc"] = mon
    lop = ClassSectionFactory(hoc_ky=hk, **kw)
    ScheduleFactory(
        lop_hoc_phan=lop, thu=thu, tiet_bat_dau=tiet, so_tiet_moi_buoi=so_tiet, phong=f"P{lop.id}"
    )
    return lop


def dang_ky(api, lop, **extra):
    return api.post(DK, {"lop_hoc_phan": lop.id}, **{**khoa_idem(), **extra})


def ma_loi(res):
    return res.data["error"]["code"]


def diem(sv, mon, diem_chu, trang_thai=TrangThaiDiem.KHOA):
    lop = ClassSectionFactory(mon_hoc=mon, trang_thai=TrangThaiLop.DA_DONG)
    return Grade.objects.create(
        lop_hoc_phan=lop, sinh_vien=sv.user, diem_chu=diem_chu, trang_thai=trang_thai
    )


# ------------------------------------------------------------------ đăng ký thành công


def test_dang_ky_thanh_cong_tang_si_so_ghi_audit_thong_bao_va_tuition_chua_tinh(sv_api, sv, hk):
    lop = tao_lop(hk)
    res = dang_ky(sv_api, lop)

    assert res.status_code == 201
    assert res.data["trang_thai"] == "DA_DANG_KY" and res.data["ma_lop"] == lop.ma_lop
    assert res.data["lich_hoc"][0]["thu"] == 2
    codes = {w["code"] for w in res.data["warnings"]}
    # chưa có bảng giá → vẫn 201 nhưng cảnh báo; dưới tối thiểu 12 TC → cảnh báo
    assert codes == {"CHUA_CAU_HINH_GIA", "DUOI_TIN_CHI_TOI_THIEU"}
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 1
    assert Tuition.objects.get(sinh_vien=sv.user, hoc_ky=hk).trang_thai == "CHUA_TINH"
    tb = Notification.objects.get()
    assert tb.audience == "USER" and tb.target_refs == [sv.user.id]
    assert AuditLog.objects.filter(doi_tuong="Registration", hanh_dong="CREATE").count() == 1


def test_hoc_phi_tinh_lai_sau_dang_ky_va_huy(sv_api, sv, hk):
    HocPhanGia.objects.create(pham_vi="MAC_DINH", don_gia=Decimal("100000"))
    mon = CourseFactory(tin_chi_ly_thuyet=2, tin_chi_thuc_hanh=1, tong_tin_chi=3)
    lop = tao_lop(hk, mon=mon)

    res = dang_ky(sv_api, lop)
    assert "CHUA_CAU_HINH_GIA" not in {w["code"] for w in res.data["warnings"]}
    tuition = Tuition.objects.get(sinh_vien=sv.user, hoc_ky=hk)
    assert (tuition.trang_thai, tuition.tong_tin_chi, tuition.tong_tien) == (
        "DA_TINH",
        3,
        Decimal("300000"),
    )

    sv_api.delete(f"{DK}{res.data['id']}/")
    tuition.refresh_from_db()
    assert (tuition.tong_tin_chi, tuition.tong_tien) == (0, Decimal("0"))


# ------------------------------------------------------------------ idempotency


def test_thieu_hoac_sai_idempotency_key_bi_chan(sv_api, hk):
    lop = tao_lop(hk)
    thieu = sv_api.post(DK, {"lop_hoc_phan": lop.id})
    assert thieu.status_code == 400 and "Idempotency-Key" in thieu.json()  # từ middleware
    sai = sv_api.post(DK, {"lop_hoc_phan": lop.id}, HTTP_IDEMPOTENCY_KEY="khong-phai-uuid")
    assert sai.status_code == 400 and "Idempotency-Key" in sai.data


def test_replay_cung_key_tra_ban_ghi_cu_khong_tang_si_so_them(sv_api, hk):
    lop = tao_lop(hk)
    key = khoa_idem()
    a = sv_api.post(DK, {"lop_hoc_phan": lop.id}, **key)
    b = sv_api.post(DK, {"lop_hoc_phan": lop.id}, **key)

    assert a.status_code == b.status_code == 201 and a.data["id"] == b.data["id"]
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 1 and Registration.objects.count() == 1
    assert AuditLog.objects.filter(doi_tuong="Registration", hanh_dong="CREATE").count() == 1


def test_key_khac_khi_da_dang_ky_tra_409_da_dang_ky(sv_api, hk):
    lop = tao_lop(hk)
    dang_ky(sv_api, lop)
    res = dang_ky(sv_api, lop)
    assert res.status_code == 409 and ma_loi(res) == "DA_DANG_KY"


def test_kich_hoat_lai_giu_nguyen_request_id_va_si_so_dung(sv_api, hk):
    lop = tao_lop(hk)
    key = khoa_idem()
    ban_dau = sv_api.post(DK, {"lop_hoc_phan": lop.id}, **key)
    sv_api.delete(f"{DK}{ban_dau.data['id']}/")
    lai = dang_ky(sv_api, lop)  # key khác

    assert lai.status_code == 201 and lai.data["id"] == ban_dau.data["id"]
    reg = Registration.objects.get()
    assert str(reg.request_id) == key["HTTP_IDEMPOTENCY_KEY"]  # không bị ghi đè
    assert reg.trang_thai == "DA_DANG_KY" and reg.ly_do_huy == "" and reg.ngay_huy is None
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 1


# ------------------------------------------------------------------ BR-01, BR-02


@pytest.mark.parametrize(
    "trang_thai", [TrangThaiLop.KE_HOACH, TrangThaiLop.DANG_HOC, TrangThaiLop.DA_DONG]
)
def test_lop_khong_o_trang_thai_mo_bi_chan(sv_api, hk, trang_thai):
    lop = tao_lop(hk, trang_thai=trang_thai)
    res = dang_ky(sv_api, lop)
    assert res.status_code == 400 and ma_loi(res) == "LOP_CHUA_MO"


def test_ngoai_cua_so_hoac_hoc_ky_dong_bi_chan(sv_api):
    now = timezone.now()
    het_han = SemesterFactory(
        cua_so_dk_bat_dau=now - timedelta(days=9), cua_so_dk_ket_thuc=now - timedelta(days=1)
    )
    res = dang_ky(sv_api, tao_lop(het_han))
    assert res.status_code == 400 and ma_loi(res) == "NGOAI_CUA_SO_DANG_KY"

    dong = SemesterFactory(
        trang_thai="DONG",
        cua_so_dk_bat_dau=now - timedelta(days=1),
        cua_so_dk_ket_thuc=now + timedelta(days=1),
    )
    assert ma_loi(dang_ky(sv_api, tao_lop(dong))) == "NGOAI_CUA_SO_DANG_KY"


def test_lop_day_tra_409(sv_api, hk):
    lop = tao_lop(hk, si_so_toi_da=1, si_so_hien_tai=1)
    res = dang_ky(sv_api, lop)
    assert res.status_code == 409 and ma_loi(res) == "LOP_DAY"


def test_lop_khong_ton_tai_va_vai_tro_khac_sv_bi_tu_choi(sv_api, login_as, hk):
    assert sv_api.post(DK, {"lop_hoc_phan": 999999}, **khoa_idem()).status_code == 400
    pdt = login_as(UserFactory(role=Role.PHONG_DAO_TAO))
    res = pdt.post(DK, {"lop_hoc_phan": tao_lop(hk).id}, **khoa_idem())
    assert res.status_code == 403 and ma_loi(res) == "KHONG_CO_QUYEN"


# ------------------------------------------------------------------ BR-03


def test_trung_lich_cung_thu_giao_tiet(sv_api, hk):
    dang_ky(sv_api, tao_lop(hk, thu=2, tiet=1, so_tiet=3))
    trung = dang_ky(sv_api, tao_lop(hk, thu=2, tiet=3, so_tiet=2))
    assert trung.status_code == 409 and ma_loi(trung) == "TRUNG_LICH"

    ke_nhau = dang_ky(sv_api, tao_lop(hk, thu=2, tiet=4, so_tiet=2))  # tiết 4–5, không giao 1–3
    khac_thu = dang_ky(sv_api, tao_lop(hk, thu=3, tiet=1, so_tiet=3))
    chu_nhat = dang_ky(sv_api, tao_lop(hk, thu=8, tiet=1, so_tiet=3))
    assert ke_nhau.status_code == khac_thu.status_code == chu_nhat.status_code == 201


def test_lop_khong_co_lich_hop_le_va_lop_hoc_ky_khac_khong_xet_trung(sv_api, hk):
    dang_ky(sv_api, tao_lop(hk, thu=2, tiet=1))
    khong_lich = ClassSectionFactory(hoc_ky=hk, trang_thai=TrangThaiLop.MO)
    assert dang_ky(sv_api, khong_lich).status_code == 201

    now = timezone.now()
    hk_khac = SemesterFactory(
        cua_so_dk_bat_dau=now - timedelta(days=1), cua_so_dk_ket_thuc=now + timedelta(days=1)
    )
    assert dang_ky(sv_api, tao_lop(hk_khac, thu=2, tiet=1)).status_code == 201


# ------------------------------------------------------------------ BR-04


def test_thieu_tien_quyet_roi_du_dieu_kien_sau_khi_dat_ke_ca_diem_d(sv_api, sv, hk):
    nen_tang, nang_cao = CourseFactory(), CourseFactory()
    Prerequisite.objects.create(mon_hoc=nang_cao, mon_tien_quyet=nen_tang)
    lop = tao_lop(hk, mon=nang_cao)

    res = dang_ky(sv_api, lop)
    assert res.status_code == 400 and ma_loi(res) == "THIEU_TIEN_QUYET"
    assert res.data["error"]["details"]["mon_tien_quyet"] == [nen_tang.ma_mon]

    diem(sv, nen_tang, "F")  # trượt tiên quyết vẫn chưa đủ
    assert ma_loi(dang_ky(sv_api, lop)) == "THIEU_TIEN_QUYET"
    diem(sv, nen_tang, "D")  # D vẫn là đạt
    assert dang_ky(sv_api, lop).status_code == 201


def test_da_dat_mon_ke_ca_d_bi_chan_con_f_hoc_lai_duoc_va_diem_chua_khoa_khong_tinh(sv_api, sv, hk):
    mon_d, mon_f, mon_nhap = CourseFactory(), CourseFactory(), CourseFactory()
    diem(sv, mon_d, "D")
    diem(sv, mon_f, "F")
    diem(sv, mon_nhap, "A", trang_thai=TrangThaiDiem.CHINH_THUC)  # chưa KHOA → chưa tính

    res = dang_ky(sv_api, tao_lop(hk, mon=mon_d))
    assert res.status_code == 400 and ma_loi(res) == "DA_DAT_MON"
    assert dang_ky(sv_api, tao_lop(hk, mon=mon_f, thu=3)).status_code == 201
    assert dang_ky(sv_api, tao_lop(hk, mon=mon_nhap, thu=4)).status_code == 201


def test_mot_mon_mot_lop_moi_hoc_ky_huy_roi_dang_ky_lop_khac_duoc(sv_api, hk):
    mon = CourseFactory()
    lop_a, lop_b = tao_lop(hk, mon=mon, thu=2), tao_lop(hk, mon=mon, thu=3)
    a = dang_ky(sv_api, lop_a)
    res = dang_ky(sv_api, lop_b)
    assert res.status_code == 409 and ma_loi(res) == "DA_DANG_KY_MON"

    sv_api.delete(f"{DK}{a.data['id']}/")
    assert dang_ky(sv_api, lop_b).status_code == 201


# ------------------------------------------------------------------ BR-07


def test_vuot_toi_da_chan_duoi_toi_thieu_chi_canh_bao(sv_api, hk):
    hk.tin_chi_toi_thieu, hk.tin_chi_toi_da, hk.tin_chi_toi_da_canh_bao = 6, 6, 3
    hk.save()
    a = dang_ky(sv_api, tao_lop(hk, thu=2))  # 3 TC < tối thiểu 6 → 201 + warning
    assert a.status_code == 201
    assert "DUOI_TIN_CHI_TOI_THIEU" in {w["code"] for w in a.data["warnings"]}
    b = dang_ky(sv_api, tao_lop(hk, thu=3))  # 6 TC = tối đa, đạt tối thiểu
    assert "DUOI_TIN_CHI_TOI_THIEU" not in {w["code"] for w in b.data["warnings"]}
    c = dang_ky(sv_api, tao_lop(hk, thu=4))
    assert c.status_code == 400 and ma_loi(c) == "VUOT_TIN_TC_TOI_DA"


def test_gpa_duoi_2_bi_gioi_han_canh_bao_gpa_null_thi_khong(login_as, hk):
    hk.tin_chi_toi_da, hk.tin_chi_toi_da_canh_bao = 24, 3
    hk.save()
    ctdt = CurriculumFactory()
    yeu = StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt, gpa=Decimal("1.50"))
    moi = StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt, gpa=None)

    api = login_as(yeu.user)
    assert dang_ky(api, tao_lop(hk, thu=2)).status_code == 201  # 3 TC = ngưỡng cảnh báo
    res = dang_ky(api, tao_lop(hk, thu=3))
    assert res.status_code == 400 and ma_loi(res) == "VUOT_TIN_TC_CANH_BAO"

    api = login_as(moi.user)  # gpa NULL (học kỳ đầu) không thuộc diện cảnh báo
    assert dang_ky(api, tao_lop(hk, thu=2)).status_code == 201
    assert dang_ky(api, tao_lop(hk, thu=3)).status_code == 201


# ------------------------------------------------------------------ hủy (BR-05)


def test_sv_huy_giam_si_so_khong_xuong_duoi_0_va_huy_hai_lan_409(sv_api, hk):
    lop = tao_lop(hk)
    reg = dang_ky(sv_api, lop)
    res = sv_api.delete(f"{DK}{reg.data['id']}/", {"ly_do": "Đổi lịch"}, format="json")

    assert res.status_code == 200 and res.data["trang_thai"] == "DA_HUY"
    assert res.data["ly_do_huy"] == "Đổi lịch"
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 0
    assert sv_api.delete(f"{DK}{reg.data['id']}/").status_code == 409

    # sĩ số bị lệch về 0 từ trước → hủy một đăng ký vẫn không âm
    reg2 = dang_ky(sv_api, lop)
    ClassSection.objects.filter(pk=lop.pk).update(si_so_hien_tai=0)
    assert sv_api.delete(f"{DK}{reg2.data['id']}/").status_code == 200
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 0


def test_sv_khong_huy_khi_lop_dang_hoc_hoac_ngoai_cua_so(sv_api, sv, hk):
    lop = tao_lop(hk)
    reg = dang_ky(sv_api, lop)
    ClassSection.objects.filter(pk=lop.pk).update(trang_thai=TrangThaiLop.DANG_HOC)
    res = sv_api.delete(f"{DK}{reg.data['id']}/")
    assert res.status_code == 409 and ma_loi(res) == "TRANG_THAI_KHONG_HOP_LE"

    ClassSection.objects.filter(pk=lop.pk).update(trang_thai=TrangThaiLop.MO)
    hk.cua_so_dk_ket_thuc = timezone.now() - timedelta(hours=1)
    hk.save()
    res = sv_api.delete(f"{DK}{reg.data['id']}/")
    assert res.status_code == 400 and ma_loi(res) == "NGOAI_CUA_SO_DANG_KY"


def test_sv_khong_huy_duoc_dang_ky_cua_nguoi_khac(sv_api, login_as, hk):
    reg = dang_ky(sv_api, tao_lop(hk))
    khac = StudentProfileFactory()
    assert login_as(khac.user).delete(f"{DK}{reg.data['id']}/").status_code == 404


def test_pdt_huy_ho_bat_buoc_ly_do_lop_dang_hoc_duoc_da_dong_thi_khong(sv_api, login_as, hk):
    pdt = login_as(UserFactory(role=Role.PHONG_DAO_TAO))
    lop = tao_lop(hk)
    reg = dang_ky(sv_api, lop)
    url = f"{DK}{reg.data['id']}/"

    thieu = pdt.delete(url, {}, format="json")
    assert thieu.status_code == 400 and "ly_do" in thieu.data

    ClassSection.objects.filter(pk=lop.pk).update(trang_thai=TrangThaiLop.DA_DONG)
    assert pdt.delete(url, {"ly_do": "x"}, format="json").status_code == 409

    ClassSection.objects.filter(pk=lop.pk).update(trang_thai=TrangThaiLop.DANG_HOC)
    ok = pdt.delete(url, {"ly_do": "Sinh viên nghỉ học"}, format="json")
    assert ok.status_code == 200 and ok.data["ly_do_huy"] == "Sinh viên nghỉ học"
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 0
    log = AuditLog.objects.filter(doi_tuong="Registration", hanh_dong="UPDATE").latest("id")
    assert log.chi_tiet["ly_do"] == "Sinh viên nghỉ học" and log.chi_tiet["boi"] == "PHONG_DAO_TAO"
    assert "Sinh viên nghỉ học" in Notification.objects.latest("id").noi_dung


# ------------------------------------------------------------------ xem


def test_lich_su_chi_cua_toi_loc_theo_trang_thai(sv_api, login_as, hk):
    a = dang_ky(sv_api, tao_lop(hk, thu=2))
    dang_ky(sv_api, tao_lop(hk, thu=3))
    sv_api.delete(f"{DK}{a.data['id']}/", {"ly_do": "Đổi ý"}, format="json")
    dang_ky(login_as(StudentProfileFactory().user), tao_lop(hk, thu=4))  # SV khác

    tat_ca = sv_api.get(f"{DK}history/")
    assert tat_ca.data["count"] == 2
    da_huy = sv_api.get(f"{DK}history/?trang_thai=DA_HUY")
    assert [r["ly_do_huy"] for r in da_huy.data["results"]] == ["Đổi ý"]
    assert sv_api.get(f"{DK}history/?hoc_ky={hk.id}").data["count"] == 2


def test_phan_quyen_danh_sach_pdt_xem_sv_khong(sv_api, login_as, hk, sv):
    dang_ky(sv_api, tao_lop(hk))
    assert sv_api.get(DK).status_code == 403
    assert login_as(UserFactory(role=Role.KE_TOAN)).get(DK).status_code == 403
    assert login_as(UserFactory(role=Role.GIAO_VIEN)).get(f"{DK}history/").status_code == 403

    pdt = login_as(UserFactory(role=Role.PHONG_DAO_TAO))
    res = pdt.get(f"{DK}?search={sv.ma_sv}")
    assert res.status_code == 200 and res.data["count"] == 1
    assert res.data["results"][0]["ma_sv"] == sv.ma_sv
    assert pdt.get(f"{DK}?nganh={sv.nganh_id}").data["count"] == 1
    assert pdt.get(f"{DK}?trang_thai=DA_HUY").data["count"] == 0


def test_xuat_excel_theo_bo_loc(sv_api, login_as, hk, sv):
    from io import BytesIO

    from openpyxl import load_workbook

    dang_ky(sv_api, tao_lop(hk))
    pdt = login_as(UserFactory(role=Role.PHONG_DAO_TAO))
    res = pdt.get(f"{DK}export/?trang_thai=DA_DANG_KY")

    assert res.status_code == 200 and "spreadsheetml" in res["Content-Type"]
    ws = load_workbook(BytesIO(res.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Mã SV" and rows[1][0] == sv.ma_sv and rows[1][7] == "Đã đăng ký"
    assert len(rows) == 2
    assert login_as(sv.user).get(f"{DK}export/").status_code == 403


# ------------------------------------------------------------------ lớp có thể đăng ký


def test_lop_co_the_dang_ky_loc_ctdt_cho_trong_trung_lich_va_meta(sv_api, sv, hk):
    def trong_ctdt(mon):
        CurriculumCourse.objects.create(ctdt=sv.ctdt, mon_hoc=mon, hoc_ky_goi_y=2)
        return mon

    hop_le = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=2, tiet=1)
    ngoai_ctdt = tao_lop(hk, thu=3)
    day = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=4, si_so_toi_da=1, si_so_hien_tai=1)
    da_dk = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=5)
    trung = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=5, tiet=2)
    ke_hoach = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=6, trang_thai=TrangThaiLop.KE_HOACH)
    dang_ky(sv_api, da_dk)

    res = sv_api.get("/api/courses/available/")
    assert res.status_code == 200
    ids = [r["id"] for r in res.data["results"]]
    assert ids == [hop_le.id]
    for loai in (ngoai_ctdt, day, da_dk, trung, ke_hoach):
        assert loai.id not in ids
    row = res.data["results"][0]
    assert row["hoc_ky_goi_y"] == 2 and row["khoi_kien_thuc"] == "chung"
    meta = res.data["meta"]
    assert meta["hoc_ky"] == hk.id and meta["trong_cua_so"] is True
    assert (meta["tin_chi_toi_thieu"], meta["tin_chi_toi_da"], meta["tin_chi_toi_da_canh_bao"]) == (
        12,
        24,
        14,
    )
    assert meta["tin_chi_da_dang_ky"] == 3 and meta["bi_canh_bao_hoc_vu"] is False


def test_available_hoc_ky_khong_hop_le_va_tim_kiem(sv_api, sv, hk):
    mon = CourseFactory(ten_mon="Trí tuệ nhân tạo")
    CurriculumCourse.objects.create(ctdt=sv.ctdt, mon_hoc=mon)
    tao_lop(hk, mon=mon)
    assert sv_api.get("/api/courses/available/?hoc_ky=abc").status_code == 400
    assert sv_api.get("/api/courses/available/?search=nhân tạo").data["count"] == 1
    assert sv_api.get("/api/courses/available/?search=khong-co").data["count"] == 0


def test_available_chi_sv(login_as):
    res = login_as(UserFactory(role=Role.GIAO_VIEN)).get("/api/courses/available/")
    assert res.status_code == 403


def test_tkb_tu_dong_tra_ca_lop_trung_lich_nhung_van_loc_ctdt_cho_trong_da_dang_ky(
    sv_api, sv, hk, login_as
):
    def trong_ctdt(mon):
        CurriculumCourse.objects.create(ctdt=sv.ctdt, mon_hoc=mon)
        return mon

    da_dk = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=5)
    trung = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=5, tiet=2)  # trùng lớp đã đăng ký
    tu_do = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=6)
    ngoai_ctdt = tao_lop(hk, thu=3)
    day = tao_lop(hk, mon=trong_ctdt(CourseFactory()), thu=4, si_so_toi_da=1, si_so_hien_tai=1)
    dang_ky(sv_api, da_dk)

    url = "/api/auto-schedule/available-courses/"
    res = sv_api.get(url)
    assert res.status_code == 200
    ids = {r["id"] for r in res.data["results"]}
    assert ids == {trung.id, tu_do.id}  # khác /courses/available/: lớp trùng lịch vẫn có mặt
    assert {ngoai_ctdt.id, day.id, da_dk.id}.isdisjoint(ids)
    thuong = {r["id"] for r in sv_api.get("/api/courses/available/").data["results"]}
    assert trung.id not in thuong  # đối chứng: endpoint thường đã lọc trùng lịch
    assert res.data["meta"]["tin_chi_da_dang_ky"] == 3  # client cần meta để tính trần tín chỉ
    assert login_as(UserFactory(role=Role.GIAO_VIEN)).get(url).status_code == 403


# ------------------------------------------------------------------ CTĐT của tôi, TKB


def test_ctdt_cua_toi_danh_dau_mon_da_dat(sv_api, sv):
    da, chua = CourseFactory(), CourseFactory()
    CurriculumCourse.objects.create(ctdt=sv.ctdt, mon_hoc=da, hoc_ky_goi_y=1)
    CurriculumCourse.objects.create(ctdt=sv.ctdt, mon_hoc=chua, hoc_ky_goi_y=2)
    diem(sv, da, "C")

    res = sv_api.get("/api/curriculums/my/")
    assert res.status_code == 200 and res.data["ma_ctdt"] == sv.ctdt.ma_ctdt
    flags = {m["ma_mon"]: m["da_dat"] for m in res.data["mon"]}
    assert flags == {da.ma_mon: True, chua.ma_mon: False}


def test_ctdt_cua_toi_404_khi_chua_gan_ctdt(login_as):
    sv = StudentProfileFactory()  # ctdt NULL
    assert login_as(sv.user).get("/api/curriculums/my/").status_code == 404


def test_thoi_khoa_bieu_sv_la_lop_da_dang_ky_gv_la_lop_phu_trach(sv_api, login_as, hk):
    lop_a, lop_b = tao_lop(hk, thu=2), tao_lop(hk, thu=3)
    dang_ky(sv_api, lop_a)
    huy = dang_ky(sv_api, lop_b)
    sv_api.delete(f"{DK}{huy.data['id']}/")  # đã hủy → không có trong TKB

    res = sv_api.get(f"/api/timetable/my/?hoc_ky={hk.id}")
    assert res.status_code == 200
    assert [x["ma_lop"] for x in res.data["lop_hoc_phan"]] == [lop_a.ma_lop]
    assert res.data["lop_hoc_phan"][0]["lich_hoc"][0]["tiet_ket_thuc"] == 3

    gv_api = login_as(lop_b.giao_vien.user)
    res = gv_api.get("/api/timetable/my/")
    assert [x["ma_lop"] for x in res.data["lop_hoc_phan"]] == [lop_b.ma_lop]
    assert login_as(UserFactory(role=Role.KE_TOAN)).get("/api/timetable/my/").status_code == 403


# ------------------------------------------------------------------ BR-06 trên PostgreSQL thật


@pytest.mark.django_db(transaction=True)
def test_hai_sv_bam_cho_cuoi_chi_mot_thanh_cong():
    """Cần PostgreSQL thật (`TEST_DB=postgres`): SQLite không có khóa dòng nên bỏ qua."""
    from concurrent.futures import ThreadPoolExecutor

    from django.db import connection, connections

    if connection.vendor != "postgresql":
        pytest.skip("Cần PostgreSQL thật để kiểm tra SELECT … FOR UPDATE")
    from rest_framework.test import APIClient
    from rest_framework_simplejwt.tokens import RefreshToken

    now = timezone.now()
    hk = SemesterFactory(
        cua_so_dk_bat_dau=now - timedelta(days=1), cua_so_dk_ket_thuc=now + timedelta(days=1)
    )
    lop = tao_lop(hk, si_so_toi_da=1)
    ctdt = CurriculumFactory()
    nguoi = [StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt).user for _ in range(2)]

    def bam(user):
        try:
            c = APIClient()
            c.credentials(HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}")
            return c.post(DK, {"lop_hoc_phan": lop.id}, **khoa_idem()).status_code
        finally:
            connections.close_all()

    with ThreadPoolExecutor(2) as ex:
        ket_qua = sorted(ex.map(bam, nguoi))

    assert ket_qua == [201, 409]
    lop.refresh_from_db()
    assert lop.si_so_hien_tai == 1
    assert Registration.objects.filter(trang_thai=TrangThaiDangKy.DA_DANG_KY).count() == 1
