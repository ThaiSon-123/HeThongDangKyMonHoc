"""Pha 9: báo cáo tài chính (KT-11) — tổng hợp, theo ngành, lọc, phân quyền, xuất Excel."""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.auth.models import Role
from apps.registration.models import Registration
from apps.tuition.models import Discount, Payment, Tuition, TuitionConfig
from core.factories import (
    ClassSectionFactory,
    MajorFactory,
    SemesterFactory,
    StudentProfileFactory,
    UserFactory,
)

pytestmark = pytest.mark.django_db
D = Decimal
BC = "/api/reports/finance/"


@pytest.fixture
def kt(login_as):
    return login_as(UserFactory(role=Role.KE_TOAN))


def hoc_phi(sv, hk, tong, thu=0, trang_thai="DA_TINH"):
    t = Tuition.objects.create(
        sinh_vien=sv.user, hoc_ky=hk, tong_tien=D(tong), tong_tin_chi=3, trang_thai=trang_thai
    )
    if thu:
        Payment.objects.create(
            tuition=t, loai="THU", so_tien=D(thu), phuong_thuc="TIEN_MAT", nguoi_thu=UserFactory()
        )
    return t


@pytest.fixture
def du_lieu():
    """CNTT: nợ 600k (quá hạn) + đủ 400k; KTPM: dư 500k (chờ hoàn); 1 SV chưa tính."""
    hk = SemesterFactory()
    TuitionConfig.objects.create(hoc_ky=hk, han_dong=timezone.localdate() - timedelta(days=3))
    cntt, ktpm = MajorFactory(ma_nganh="CNTT"), MajorFactory(ma_nganh="KTPM")
    no, du, thua, moi = (StudentProfileFactory(nganh=n) for n in (cntt, cntt, ktpm, ktpm))
    hoc_phi(no, hk, 1_000_000, thu=400_000)
    hoc_phi(du, hk, 400_000, thu=400_000)
    hoc_phi(thua, hk, 500_000, thu=1_000_000)
    hoc_phi(moi, hk, 0, trang_thai="CHUA_TINH")
    return {"hk": hk, "cntt": cntt, "ktpm": ktpm, "no": no, "thua": thua}


def test_tong_hop_khop_tung_sinh_vien(kt, du_lieu):
    res = kt.get(BC, {"hoc_ky": du_lieu["hk"].id})
    assert res.status_code == 200
    t = res.data["tong"]
    assert t["phai_thu"] == "1900000.00"  # tiền là chuỗi, như các API học phí khác
    assert (D(t["phai_thu"]), D(t["da_thu"]), D(t["con_no"]), D(t["cho_hoan"])) == (
        D("1900000"),
        D("1800000"),
        D("600000"),
        D("500000"),
    )
    assert (t["so_sinh_vien"], t["so_chua_tinh"], t["so_con_no"]) == (4, 1, 1)
    assert (t["so_qua_han"], t["so_cho_hoan"]) == (1, 1)
    # ràng buộc kế toán: phải thu − đã thu = còn nợ − chờ hoàn
    assert D(t["phai_thu"]) - D(t["da_thu"]) == D(t["con_no"]) - D(t["cho_hoan"])
    nganh = {n["nganh_ma"]: n for n in res.data["theo_nganh"]}
    assert (D(nganh["CNTT"]["phai_thu"]), D(nganh["CNTT"]["da_thu"])) == (
        D("1400000"),
        D("800000"),
    )
    assert (D(nganh["KTPM"]["phai_thu"]), D(nganh["KTPM"]["da_thu"])) == (
        D("500000"),
        D("1000000"),
    )
    assert res.data["theo_nganh"][0]["nganh_ma"] == "CNTT"  # phải thu lớn nhất trước


def test_mien_giam_tru_vao_phai_thu(kt, du_lieu):
    Discount.objects.create(
        sinh_vien=du_lieu["no"].user, hoc_ky=du_lieu["hk"], loai="MIEN_GIAM", so_tien=D("200000")
    )
    res = kt.get(BC, {"hoc_ky": du_lieu["hk"].id})
    assert D(res.data["tong"]["phai_thu"]) == D("1700000")
    assert D(res.data["tong"]["con_no"]) == D("400000")


def test_loc_nganh_lop_va_sinh_vien(kt, du_lieu):
    hk = du_lieu["hk"]
    tham_so = {"hoc_ky": hk.id}
    res = kt.get(BC, {**tham_so, "nganh": du_lieu["ktpm"].id})
    assert res.data["tong"]["so_sinh_vien"] == 2
    assert D(res.data["tong"]["cho_hoan"]) == D("500000")
    res = kt.get(BC, {**tham_so, "sinh_vien": du_lieu["no"].user_id})
    assert res.data["tong"]["so_sinh_vien"] == 1
    assert D(res.data["tong"]["con_no"]) == D("600000")
    lop = ClassSectionFactory(hoc_ky=hk)
    Registration.objects.create(sinh_vien=du_lieu["thua"].user, lop_hoc_phan=lop)
    res = kt.get(BC, {**tham_so, "lop_hoc_phan": lop.id})
    assert res.data["tong"]["so_sinh_vien"] == 1
    # danh sách học phí (bảng chi tiết của màn báo cáo) cũng lọc được theo sinh viên
    rows = kt.get("/api/tuition/", {**tham_so, "sinh_vien": du_lieu["no"].user_id}).data["results"]
    assert [r["ma_sv"] for r in rows] == [du_lieu["no"].ma_sv]
    assert rows[0]["nganh_ma"] == "CNTT"  # cột Ngành của bảng chi tiết


def test_hoc_ky_rong_cho_so_0(kt):
    hk = SemesterFactory()
    res = kt.get(BC, {"hoc_ky": hk.id})
    assert D(res.data["tong"]["phai_thu"]) == 0 and res.data["theo_nganh"] == []


def test_tham_so_sai_va_phan_quyen(login_as, kt):
    assert kt.get(BC, {"nganh": "x"}).status_code == 400
    assert kt.get(BC, {"sinh_vien": "x"}).status_code == 400
    assert kt.get("/api/tuition/", {"lop_hoc_phan": "x"}).status_code == 400
    for role in (Role.PHONG_DAO_TAO, Role.SINH_VIEN, Role.GIAO_VIEN, Role.ADMIN):
        c = login_as(UserFactory(role=role))
        assert c.get(BC).status_code == 403
        assert c.get(BC + "export/").status_code == 403


def test_xuat_excel_bao_cao_tai_chinh(kt, du_lieu):
    res = kt.get(BC + "export/", {"hoc_ky": du_lieu["hk"].id})
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Tổng hợp", "Theo ngành", "Chi tiết theo sinh viên"]
    tong = {r[0]: r[1] for r in wb["Tổng hợp"].iter_rows(min_row=2, values_only=True)}
    assert (tong["Phải thu"], tong["Đã thu"], tong["Còn nợ"], tong["Chờ hoàn tiền"]) == (
        1_900_000,
        1_800_000,
        600_000,
        500_000,
    )
    chi_tiet = list(wb["Chi tiết theo sinh viên"].iter_rows(min_row=2, values_only=True))
    assert len(chi_tiet) == 4
    trang_thai = {r[0]: r[9] for r in chi_tiet}
    assert trang_thai[du_lieu["no"].ma_sv] == "Còn nợ (quá hạn)"
    assert trang_thai[du_lieu["thua"].ma_sv] == "Chờ hoàn tiền"
