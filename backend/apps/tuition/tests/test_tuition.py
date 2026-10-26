"""Học phí (BR-15..18, BR-26): bảng giá, miễn giảm, tính, thu/hoàn, công nợ, biên lai, tuition/my, job quá hạn."""

import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.auth.models import Role
from apps.config.models import SystemConfig
from apps.notifications.models import Notification
from apps.registration.models import Registration
from apps.training.models import CurriculumCourse, TrangThaiLop
from apps.tuition.models import Discount, HocPhanGia, Payment, Tuition, TuitionConfig
from apps.tuition.services import gia_tri_dan_xuat, thong_bao_qua_han, tinh_hoc_phi
from core.factories import (
    ClassSectionFactory,
    CourseFactory,
    CurriculumFactory,
    SemesterFactory,
    StudentProfileFactory,
    UserFactory,
)

pytestmark = pytest.mark.django_db
D = Decimal
PAY = "/api/tuition/payments/"


@pytest.fixture
def kt(login_as):
    return login_as(UserFactory(role=Role.KE_TOAN))


def khoa_idem():
    return {"HTTP_IDEMPOTENCY_KEY": str(uuid.uuid4())}


def gia(pham_vi, don_gia, tham_chieu=None, hoc_ky=None):
    return HocPhanGia.objects.create(
        pham_vi=pham_vi, tham_chieu=tham_chieu, don_gia=D(don_gia), hoc_ky=hoc_ky
    )


def dang_ky(sv, lop):
    return Registration.objects.create(sinh_vien=sv.user, lop_hoc_phan=lop)


def hoc_phi(sv, hk, tong, trang_thai="DA_TINH"):
    return Tuition.objects.create(
        sinh_vien=sv.user, hoc_ky=hk, tong_tien=D(tong), tong_tin_chi=3, trang_thai=trang_thai
    )


# --------------------------------------------------------------- giá trị dẫn xuất (mục 5)


class _D:  # discount giả để test hàm thuần
    def __init__(self, so_tien=None, ty_le=None):
        self.so_tien, self.ty_le = so_tien, ty_le


class _P:
    def __init__(self, loai, so_tien):
        self.loai, self.so_tien = loai, D(so_tien)


class _T:
    def __init__(self, tong, trang_thai="DA_TINH"):
        self.tong_tien, self.trang_thai = D(tong), trang_thai


def test_dan_xuat_mien_giam_toi_da_bang_tong_tien_va_ty_le():
    dx = gia_tri_dan_xuat(_T(1000), [_D(so_tien=D(300)), _D(ty_le=D(10))], [])
    assert (dx["mien_giam"], dx["phai_nop"], dx["con_no"]) == (
        D("400.00"),
        D("600.00"),
        D("600.00"),
    )
    assert gia_tri_dan_xuat(_T(1000), [_D(so_tien=D(5000))], [])["mien_giam"] == D(
        "1000.00"
    )  # không vượt tổng


def test_dan_xuat_thu_hoan_con_no_can_hoan_va_hien_thi():
    t = _T(1000)
    assert gia_tri_dan_xuat(t, [], [])["hien_thi"] == "CON_NO"
    assert gia_tri_dan_xuat(t, [], [_P("THU", 1000)])["hien_thi"] == "DA_DONG"
    qua = gia_tri_dan_xuat(t, [], [_P("THU", 1200)])
    assert (qua["can_hoan"], qua["con_no"], qua["hien_thi"]) == (D(200), D(0), "CHO_HOAN_TIEN")
    sau_hoan = gia_tri_dan_xuat(t, [], [_P("THU", 1200), _P("HOAN", 200)])
    assert (sau_hoan["thuc_thu"], sau_hoan["hien_thi"]) == (D(1000), "DA_DONG")
    assert gia_tri_dan_xuat(_T(0, "CHUA_TINH"), [], [])["hien_thi"] == "CHUA_TINH"


# --------------------------------------------------------------- tính học phí (BR-15..17)


def test_tinh_hoc_phi_theo_bang_gia_lt_gia_goc_th_gia_th():
    hk = SemesterFactory()
    ctdt = CurriculumFactory()
    sv = StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt)
    mon = CourseFactory(tin_chi_ly_thuyet=2, tin_chi_thuc_hanh=1, tong_tin_chi=3)
    CurriculumCourse.objects.create(ctdt=ctdt, mon_hoc=mon, khoi_kien_thuc="chuyen_nganh")
    dang_ky(sv, ClassSectionFactory(hoc_ky=hk, mon_hoc=mon, trang_thai=TrangThaiLop.MO))
    gia("MAC_DINH", 100)
    gia("KHOI_KIEN_THUC", 150, "chuyen_nganh")
    gia("GIA_TH", 70)

    t = tinh_hoc_phi(sv.user, hk)
    assert (t.trang_thai, t.tong_tin_chi, t.tong_tien) == ("DA_TINH", 3, D("370"))  # 2×150 + 1×70


def test_truc_hoc_ky_thang_truc_mac_dinh_va_thieu_gia_la_loi_400(kt):
    hk = SemesterFactory()
    sv = StudentProfileFactory()
    mon = CourseFactory(tin_chi_ly_thuyet=3, tong_tin_chi=3)
    dang_ky(sv, ClassSectionFactory(hoc_ky=hk, mon_hoc=mon, trang_thai=TrangThaiLop.MO))

    thieu = kt.post(f"/api/tuition/calculate/?hoc_ky={hk.id}")
    assert thieu.status_code == 400 and thieu.data["error"]["code"] == "CHUA_CAU_HINH_GIA"
    assert not Tuition.objects.filter(trang_thai="DA_TINH").exists()  # cả lô bị hủy

    gia("MAC_DINH", 100)
    gia("MAC_DINH", 200, hoc_ky=hk)  # bản ghi của học kỳ thắng bản mặc định
    ok = kt.post(f"/api/tuition/calculate/?hoc_ky={hk.id}")
    assert ok.status_code == 200 and ok.data["so_sinh_vien"] == 1
    assert Tuition.objects.get(sinh_vien=sv.user).tong_tien == D("600")
    assert AuditLog.objects.filter(
        doi_tuong="Tuition", chi_tiet__hanh_dong="tinh_hang_loat"
    ).exists()


def test_calculate_chi_ke_toan(login_as):
    for r in (Role.SINH_VIEN, Role.PHONG_DAO_TAO, Role.ADMIN, Role.GIAO_VIEN):
        assert login_as(UserFactory(role=r)).post("/api/tuition/calculate/").status_code == 403


def test_huy_mon_da_thu_du_phat_sinh_can_hoan_roi_hoan_xong_dong_du(kt):
    from apps.tuition.services import dan_xuat_cua

    hk = SemesterFactory()
    ctdt = CurriculumFactory()
    sv = StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt)
    mon = CourseFactory(tin_chi_ly_thuyet=3, tong_tin_chi=3)
    reg = dang_ky(sv, ClassSectionFactory(hoc_ky=hk, mon_hoc=mon, trang_thai=TrangThaiLop.MO))
    gia("MAC_DINH", 100_000)
    t = tinh_hoc_phi(sv.user, hk)
    assert thu(kt, t, 300_000).status_code == 201  # đóng đủ

    reg.trang_thai = "DA_HUY"  # hủy môn → tính lại: phải nộp 0
    reg.save()
    t = tinh_hoc_phi(sv.user, hk)
    dx = dan_xuat_cua(t)
    assert (dx["phai_nop"], dx["con_no"], dx["can_hoan"], dx["hien_thi"]) == (
        D(0),
        D(0),
        D(300_000),
        "CHO_HOAN_TIEN",
    )

    assert thu(kt, t, 300_001, loai="HOAN").status_code == 400
    assert thu(kt, t, 300_000, loai="HOAN").status_code == 201
    assert dan_xuat_cua(t)["hien_thi"] == "DA_DONG"


# --------------------------------------------------------------- bảng giá / cấu hình / miễn giảm


def test_bang_gia_validate_pham_vi_tham_chieu_va_unique_mac_dinh(kt):
    hk = SemesterFactory()
    body = {"hoc_ky": hk.id, "pham_vi": "MAC_DINH", "don_gia": "100000"}
    assert kt.post("/api/tuition/pricing/", body).status_code == 201
    trung = kt.post(
        "/api/tuition/pricing/", body
    )  # UNIQUE (hoc_ky, MAC_DINH) kể cả tham_chieu NULL
    assert trung.status_code == 400 and "pham_vi" in trung.data
    assert (
        kt.post("/api/tuition/pricing/", {"pham_vi": "MAC_DINH", "don_gia": "90000"}).status_code
        == 201
    )
    assert (
        kt.post("/api/tuition/pricing/", {**body, "tham_chieu": "X"}).status_code == 400
    )  # MAC_DINH không tham chiếu
    assert (
        kt.post("/api/tuition/pricing/", {"pham_vi": "NGANH", "don_gia": "1"}).status_code == 400
    )  # thiếu mã ngành
    assert (
        kt.post(
            "/api/tuition/pricing/", {"pham_vi": "NGANH", "tham_chieu": "KHONGCO", "don_gia": "1"}
        ).status_code
        == 400
    )
    assert (
        kt.post(
            "/api/tuition/pricing/",
            {"pham_vi": "KHOI_KIEN_THUC", "tham_chieu": "la", "don_gia": "1"},
        ).status_code
        == 400
    )
    assert (
        kt.post(
            "/api/tuition/pricing/",
            {"pham_vi": "KHOI_KIEN_THUC", "tham_chieu": "chung", "don_gia": "-5"},
        ).status_code
        == 400
    )
    ok = kt.post(
        "/api/tuition/pricing/",
        {"pham_vi": "KHOI_KIEN_THUC", "tham_chieu": "chung", "don_gia": "120000"},
    )
    assert ok.status_code == 201
    assert kt.delete(f"/api/tuition/pricing/{ok.data['id']}/").status_code == 204
    assert AuditLog.objects.filter(doi_tuong="HocPhanGia", hanh_dong="CREATE").exists()


def test_cau_hinh_han_dong_moi_hoc_ky_mot_ban_ghi(kt):
    hk = SemesterFactory()
    assert (
        kt.post("/api/tuition/config/", {"hoc_ky": hk.id, "han_dong": "2027-03-01"}).status_code
        == 201
    )
    assert (
        kt.post("/api/tuition/config/", {"hoc_ky": hk.id, "han_dong": "2027-04-01"}).status_code
        == 400
    )
    cid = TuitionConfig.objects.get().id
    assert kt.patch(f"/api/tuition/config/{cid}/", {"han_dong": "2027-04-01"}).status_code == 200


def test_mien_giam_dung_mot_hinh_thuc_chi_cho_sinh_vien_va_ghi_audit(kt):
    hk = SemesterFactory()
    sv = StudentProfileFactory()
    base = {"sinh_vien": sv.user_id, "hoc_ky": hk.id, "loai": "MIEN_GIAM", "ly_do": "Hộ nghèo"}
    assert kt.post("/api/tuition/discounts/", {**base}).status_code == 400  # thiếu cả hai
    assert (
        kt.post("/api/tuition/discounts/", {**base, "so_tien": "1000", "ty_le": "10"}).status_code
        == 400
    )
    assert kt.post("/api/tuition/discounts/", {**base, "ty_le": "120"}).status_code == 400
    assert kt.post("/api/tuition/discounts/", {**base, "so_tien": "0"}).status_code == 400
    assert (
        kt.post(
            "/api/tuition/discounts/",
            {**base, "sinh_vien": UserFactory(role=Role.GIAO_VIEN).id, "ty_le": "10"},
        ).status_code
        == 400
    )
    ok = kt.post("/api/tuition/discounts/", {**base, "ty_le": "25"})
    assert ok.status_code == 201 and ok.data["ma_sv"] == sv.ma_sv
    sua = kt.patch(
        f"/api/tuition/discounts/{ok.data['id']}/",
        {"ty_le": None, "so_tien": "500000"},
        format="json",
    )
    assert sua.status_code == 200
    assert AuditLog.objects.filter(doi_tuong="Discount").count() == 2  # BR-18


def test_tra_sv_cho_ke_toan_chi_thong_tin_toi_thieu(kt):
    sv = StudentProfileFactory()
    res = kt.get(f"/api/tuition/students/?search={sv.ma_sv}")
    assert res.data["count"] == 1
    assert set(res.data["results"][0]) == {
        "id",
        "ma_sv",
        "ho_ten",
        "nganh_ten",
    }  # không lộ email/sdt/gpa


# --------------------------------------------------------------- thu / hoàn (BR-26)


def thu(api, t, so_tien, loai="THU", **extra):
    return api.post(
        PAY,
        {"tuition": t.id, "loai": loai, "so_tien": str(so_tien), "phuong_thuc": "TIEN_MAT"},
        **extra,
    )


def test_thu_khong_vuot_con_no_hoan_khong_vuot_can_hoan(kt):
    hk = SemesterFactory()
    sv = StudentProfileFactory()
    t = hoc_phi(sv, hk, 1_000_000)
    assert thu(kt, t, 400_000).status_code == 201
    qua = thu(kt, t, 700_000)
    assert qua.status_code == 400 and qua.data["error"]["code"] == "SO_TIEN_KHONG_HOP_LE"
    assert thu(kt, t, 600_000).status_code == 201  # đủ

    hoan_qua = thu(kt, t, 1, loai="HOAN")  # chưa có gì để hoàn
    assert hoan_qua.status_code == 400 and hoan_qua.data["error"]["code"] == "SO_TIEN_KHONG_HOP_LE"

    Discount.objects.create(
        sinh_vien=sv.user, hoc_ky=hk, loai="MIEN_GIAM", so_tien=D(100_000)
    )  # phát sinh dư
    assert thu(kt, t, 100_001, loai="HOAN").status_code == 400
    ok = thu(kt, t, 100_000, loai="HOAN")
    assert ok.status_code == 201 and ok.data["bien_lai_ma"].startswith("BL")
    chi_tiet = kt.get(f"/api/tuition/{t.id}/")
    assert chi_tiet.data["hien_thi"] == "DA_DONG" and D(chi_tiet.data["thuc_thu"]) == D("900000.00")
    assert len(chi_tiet.data["payments"]) == 3


def test_so_tien_khong_hop_le_va_hoc_phi_chua_tinh(kt):
    hk = SemesterFactory()
    sv = StudentProfileFactory()
    t = hoc_phi(sv, hk, 1000)
    assert thu(kt, t, 0).status_code == 400
    assert thu(kt, t, -5).status_code == 400
    chua = hoc_phi(StudentProfileFactory(), hk, 1000, trang_thai="CHUA_TINH")
    res = thu(kt, chua, 500)
    assert res.status_code == 409 and res.data["error"]["code"] == "TRANG_THAI_KHONG_HOP_LE"
    assert thu(kt, t, 10).data["tuition"] == t.id
    assert (
        kt.post(
            PAY, {"tuition": 999999, "loai": "THU", "so_tien": "1", "phuong_thuc": "KHAC"}
        ).status_code
        == 400
    )


def test_replay_idempotency_key_va_audit_mot_lan(kt):
    t = hoc_phi(StudentProfileFactory(), SemesterFactory(), 1_000_000)
    key = khoa_idem()
    a = thu(kt, t, 300_000, **key)
    b = thu(kt, t, 300_000, **key)
    assert a.status_code == b.status_code == 201 and a.data["id"] == b.data["id"]
    assert Payment.objects.count() == 1
    assert AuditLog.objects.filter(doi_tuong="Payment", hanh_dong="CREATE").count() == 1
    assert thu(kt, t, 1, HTTP_IDEMPOTENCY_KEY="khong-phai-uuid").status_code == 400
    assert thu(kt, t, 1).status_code == 201  # không bắt buộc key


def test_phan_quyen_thanh_toan(login_as):
    t = hoc_phi(StudentProfileFactory(), SemesterFactory(), 1000)
    for r in (Role.SINH_VIEN, Role.PHONG_DAO_TAO, Role.ADMIN, Role.GIAO_VIEN):
        assert thu(login_as(UserFactory(role=r)), t, 10).status_code == 403


@pytest.mark.django_db(transaction=True)
def test_thanh_toan_dong_thoi_khong_vuot_phai_nop():
    """Cần PostgreSQL thật: hai KT cùng thu hết công nợ → chỉ một thành công."""
    from concurrent.futures import ThreadPoolExecutor

    from django.db import connection, connections
    from rest_framework.test import APIClient
    from rest_framework_simplejwt.tokens import RefreshToken

    if connection.vendor != "postgresql":
        pytest.skip("Cần PostgreSQL thật để kiểm tra SELECT … FOR UPDATE")
    t = hoc_phi(StudentProfileFactory(), SemesterFactory(), 1_000_000)
    kts = [UserFactory(role=Role.KE_TOAN) for _ in range(2)]

    def bam(u):
        try:
            c = APIClient()
            c.credentials(HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(u).access_token}")
            return thu(c, t, 1_000_000).status_code
        finally:
            connections.close_all()

    with ThreadPoolExecutor(2) as ex:
        ket_qua = sorted(ex.map(bam, kts))
    assert ket_qua == [201, 400]
    assert Payment.objects.count() == 1


# --------------------------------------------------------------- danh sách, công nợ, biên lai


def test_danh_sach_hoc_phi_loc_va_gia_tri_dan_xuat(kt):
    hk = SemesterFactory()
    cfg = TuitionConfig.objects.create(hoc_ky=hk, han_dong=timezone.localdate() - timedelta(days=1))
    a, b, c = StudentProfileFactory(), StudentProfileFactory(), StudentProfileFactory()
    ta, tb, tc = hoc_phi(a, hk, 1000), hoc_phi(b, hk, 1000), hoc_phi(c, hk, 1000, "CHUA_TINH")
    Payment.objects.create(
        tuition=tb, loai="THU", so_tien=D(1000), phuong_thuc="KHAC", nguoi_thu=kt_user(kt)
    )
    Payment.objects.create(
        tuition=tc, loai="THU", so_tien=D(1), phuong_thuc="KHAC", nguoi_thu=kt_user(kt)
    )

    res = kt.get(f"/api/tuition/?hoc_ky={hk.id}")
    assert res.status_code == 200 and res.data["count"] == 3
    by = {r["ma_sv"]: r for r in res.data["results"]}
    assert by[a.ma_sv]["hien_thi"] == "CON_NO" and by[a.ma_sv]["qua_han"] is True
    assert by[b.ma_sv]["hien_thi"] == "DA_DONG" and by[b.ma_sv]["qua_han"] is False
    assert by[c.ma_sv]["hien_thi"] == "CHUA_TINH"
    assert set(by[a.ma_sv]) >= {
        "mien_giam",
        "phai_nop",
        "thuc_thu",
        "con_no",
        "can_hoan",
        "han_dong",
    }
    assert "email" not in by[a.ma_sv] and "sdt" not in by[a.ma_sv]

    assert [
        r["ma_sv"] for r in kt.get(f"/api/tuition/?hoc_ky={hk.id}&hien_thi=CON_NO").data["results"]
    ] == [a.ma_sv]
    assert kt.get(f"/api/tuition/?hoc_ky={hk.id}&search={b.ma_sv}").data["count"] == 1
    assert kt.get(f"/api/tuition/?hoc_ky={hk.id}&nganh={a.nganh_id}").data["count"] == 1
    assert cfg and ta and tb and tc


def kt_user(api):
    from apps.auth.models import User

    return User.objects.filter(role=Role.KE_TOAN).first()


def test_cong_no_gom_con_no_qua_han_va_cho_hoan_tien(kt):
    hk = SemesterFactory()
    TuitionConfig.objects.create(hoc_ky=hk, han_dong=timezone.localdate() - timedelta(days=3))
    no, du, xong = StudentProfileFactory(), StudentProfileFactory(), StudentProfileFactory()
    hoc_phi(no, hk, 1000)
    t_du = hoc_phi(du, hk, 1000)
    t_xong = hoc_phi(xong, hk, 1000)
    u = kt_user(kt) or UserFactory(role=Role.KE_TOAN)
    Payment.objects.create(
        tuition=t_du, loai="THU", so_tien=D(1500), phuong_thuc="KHAC", nguoi_thu=u
    )
    Payment.objects.create(
        tuition=t_xong, loai="THU", so_tien=D(1000), phuong_thuc="KHAC", nguoi_thu=u
    )

    res = kt.get(f"/api/tuition/debts/?hoc_ky={hk.id}")
    by = {r["ma_sv"]: r for r in res.data["results"]}
    assert set(by) == {no.ma_sv, du.ma_sv}  # đã đóng đủ thì không có trong công nợ
    assert by[no.ma_sv]["qua_han"] is True and by[du.ma_sv]["hien_thi"] == "CHO_HOAN_TIEN"
    assert D(by[du.ma_sv]["can_hoan"]) == D("500.00")
    assert kt.get(f"/api/tuition/debts/?hoc_ky={hk.id}&hien_thi=CHO_HOAN_TIEN").data["count"] == 1
    assert (
        kt.get(f"/api/tuition/debts/?hoc_ky={hk.id}&qua_han=1").data["count"] == 1
    )  # chỉ SV còn nợ quá hạn


def test_bien_lai_excel(kt):
    t = hoc_phi(StudentProfileFactory(), SemesterFactory(), 1_000_000)
    p = thu(kt, t, 250_000).data
    res = kt.get(f"{PAY}{p['id']}/receipt/")
    assert res.status_code == 200 and "spreadsheetml" in res["Content-Type"]
    rows = dict(list(load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True))[1:])
    assert (
        rows["Mã biên lai"] == p["bien_lai_ma"]
        and rows["Số tiền"] == 250000
        and rows["Loại"] == "Thu"
    )


# --------------------------------------------------------------- tuition/my (SV-12)


def test_tuition_my_chi_tiet_giao_dich_va_thong_tin_chuyen_khoan(login_as, kt):
    hk = SemesterFactory()
    ctdt = CurriculumFactory()
    sv = StudentProfileFactory(nganh=ctdt.nganh, ctdt=ctdt)
    mon = CourseFactory(tin_chi_ly_thuyet=3, tong_tin_chi=3)
    dang_ky(sv, ClassSectionFactory(hoc_ky=hk, mon_hoc=mon, trang_thai=TrangThaiLop.MO))
    gia("MAC_DINH", 100000)
    tinh_hoc_phi(sv.user, hk)
    for khoa, v in (
        ("ck_ngan_hang", "Vietcombank"),
        ("ck_so_tai_khoan", "0123456789"),
        ("ck_chu_tai_khoan", "TRUONG DAI HOC X"),
    ):
        SystemConfig.objects.create(khoa=khoa, gia_tri=v)
    Discount.objects.create(sinh_vien=sv.user, hoc_ky=hk, loai="HOC_BONG", ty_le=D(10))
    Payment.objects.create(
        tuition=Tuition.objects.get(),
        loai="THU",
        so_tien=D(100000),
        phuong_thuc="CHUYEN_KHOAN",
        nguoi_thu=kt_user(kt) or UserFactory(role=Role.KE_TOAN),
    )

    res = login_as(sv.user).get(f"/api/tuition/my/?hoc_ky={hk.id}")
    assert res.status_code == 200
    d = res.data
    assert (
        D(d["tong_tien"]),
        D(d["mien_giam"]),
        D(d["phai_nop"]),
        D(d["thuc_thu"]),
        D(d["con_no"]),
    ) == (
        D("300000"),
        D("30000"),
        D("270000"),
        D("100000"),
        D("170000"),
    )
    assert d["hien_thi"] == "CON_NO"
    assert [h["id"] for h in d["cac_hoc_ky"]] == [hk.id]  # kỳ đang xem có trong danh sách chọn kỳ
    assert d["chi_tiet"][0]["ma_mon"] == mon.ma_mon and D(d["chi_tiet"][0]["thanh_tien"]) == D(
        "300000"
    )
    assert len(d["payments"]) == 1
    tt = d["thanh_toan"]
    assert tt["ngan_hang"] == "Vietcombank" and tt["so_tai_khoan"] == "0123456789"
    assert tt["noi_dung_ck"] == f"HOCPHI {sv.ma_sv} HK{hk.id}"


def test_tuition_my_chua_co_dang_ky_va_phan_quyen(login_as):
    hk = SemesterFactory()
    sv = StudentProfileFactory()
    res = login_as(sv.user).get(f"/api/tuition/my/?hoc_ky={hk.id}")
    assert (
        res.status_code == 200
        and res.data["hien_thi"] == "CHUA_TINH"
        and res.data["chi_tiet"] == []
    )
    assert D(res.data["tong_tien"]) == D(0)
    for r in (Role.GIAO_VIEN, Role.KE_TOAN, Role.PHONG_DAO_TAO):
        assert login_as(UserFactory(role=r)).get("/api/tuition/my/").status_code == 403


def test_sv_khong_xem_hoc_phi_cua_nguoi_khac(login_as):
    hk = SemesterFactory()
    a, b = StudentProfileFactory(), StudentProfileFactory()
    hoc_phi(b, hk, 5000)
    res = login_as(a.user).get(f"/api/tuition/my/?hoc_ky={hk.id}")
    assert D(res.data["tong_tien"]) == 0 and res.data["ma_sv"] == a.ma_sv
    assert login_as(a.user).get("/api/tuition/").status_code == 403


# --------------------------------------------------------------- job quá hạn


def test_job_qua_han_bao_sv_mot_lan_va_ke_toan_khi_co_sv_moi():
    hk = SemesterFactory()
    TuitionConfig.objects.create(hoc_ky=hk, han_dong=date(2026, 3, 1))
    no, du, chua_het_han = StudentProfileFactory(), StudentProfileFactory(), StudentProfileFactory()
    hoc_phi(no, hk, 1000)
    t_du = hoc_phi(du, hk, 1000)
    u = UserFactory(role=Role.KE_TOAN)
    Payment.objects.create(
        tuition=t_du, loai="THU", so_tien=D(1000), phuong_thuc="KHAC", nguoi_thu=u
    )
    hk2 = SemesterFactory()
    TuitionConfig.objects.create(hoc_ky=hk2, han_dong=date(2099, 1, 1))
    hoc_phi(chua_het_han, hk2, 1000)

    assert thong_bao_qua_han(hom_nay=date(2026, 3, 2)) == 1  # chỉ SV còn nợ + quá hạn
    sv_tb = Notification.objects.filter(target_refs=[no.user_id])
    assert sv_tb.count() == 1 and "quá hạn" in sv_tb.get().tieu_de
    assert Notification.objects.filter(target_refs=[u.id]).count() == 1  # tổng hợp cho KT
    assert not Notification.objects.filter(target_refs=[du.user_id]).exists()

    assert thong_bao_qua_han(hom_nay=date(2026, 3, 3)) == 0  # chạy lại: không báo trùng
    assert Notification.objects.count() == 2


def test_task_celery_goi_service():
    from jobs.tasks import tuition_overdue_notify

    assert tuition_overdue_notify() == 0


def test_bang_gia_loc_ap_dung_cho_gom_gia_moi_hoc_ky(kt):
    hk1, hk2 = SemesterFactory(), SemesterFactory()
    gia("MAC_DINH", 100)  # mọi học kỳ
    gia("MAC_DINH", 200, hoc_ky=hk1)
    gia("MAC_DINH", 300, hoc_ky=hk2)

    def don_gia(query):
        res = kt.get(f"/api/tuition/pricing/?{query}")
        assert res.status_code == 200, res.data
        return sorted(D(r["don_gia"]) for r in res.json()["results"])

    assert don_gia(f"ap_dung_cho={hk1.id}") == [100, 200]
    assert don_gia(f"ap_dung_cho={hk2.id}") == [100, 300]
    assert don_gia(f"hoc_ky={hk1.id}") == [200]  # tham số cũ: đúng kỳ, không kèm giá chung
    assert don_gia("") == [100, 200, 300]
    assert kt.get("/api/tuition/pricing/?ap_dung_cho=abc").status_code == 400
