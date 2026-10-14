"""Xuất Excel (openpyxl) — plan.md mục 8.2. Mọi file xuất đi qua đây: dòng tiêu đề in đậm, cột vừa nội dung."""

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_RONG = 60


def _them_sheet(ws, headers, rows):
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    for i, col in enumerate(ws.columns, 1):
        rong = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(rong + 2, MAX_RONG)


def xlsx_nhieu_sheet(filename, sheets):
    """`sheets`: danh sách `(tên_sheet, headers, rows)` → một file Excel nhiều sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    for ten, headers, rows in sheets:
        _them_sheet(wb.create_sheet(ten[:31]), headers, rows)  # Excel giới hạn tên sheet 31 ký tự
    buf = BytesIO()
    wb.save(buf)
    res = HttpResponse(buf.getvalue(), content_type=XLSX)
    res["Content-Disposition"] = f'attachment; filename="{filename}"'
    return res


def xlsx_response(filename, headers, rows, sheet="Dữ liệu"):
    return xlsx_nhieu_sheet(filename, [(sheet, headers, rows)])
